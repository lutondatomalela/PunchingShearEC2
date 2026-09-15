"""Exportações exclusivamente a partir de uma cópia do resultado da execução."""
import json
import math
from pathlib import Path
from .version import REPOSITORY

INPUT_LABELS={
    'project':'Projeto','support':'Apoio / laje','combination':'Combinação ELU',
    'longitudinal_layout':'Definição automática da armadura longitudinal',
    'laje_d':'d médio (m)','laje_dx':'dx (m)','laje_dy':'dy (m)',
    'betão_fck':'fck (MPa)','aço_fyk':'fyk (MPa)','aço_fywk':'fywk (MPa)',
    'pilar_tipo':'Posição do pilar','pilar_forma':'Forma do pilar','pilar_c1':'c1 / D (m)',
    'pilar_c2':'c2 (m)','edge_distance_m':'Face do pilar ao bordo g (m)','V_Ed':'VEd (N)','M_Edx':'MEdx (N.m)','M_Edy':'MEdy (N.m)',
    'sigma_cp':'Compressão média no plano (MPa)','is_sapata':'Sapata',
    'sigma_gd_kpa':'Pressão líquida indicada (kPa; 0 = automática)','u1_ineffective':'Dedução isolada de u1 (m; desativada)',
    'gamma_C':'gamma_C','gamma_S':'gamma_S','beta_mode':'Modo de beta',
    'beta_manual':'Beta fundamentado pelo projetista','beta_reference':'Referência da análise de beta',
    'interior_beta_method':'Expressão de beta para pilar interior retangular',
    'allow_biaxial_envelope':'Combinação conservadora de (6.39) nas duas direções selecionada',
    'laje_As_lx_cm2pm':'As,x média (cm²/m)','laje_As_ly_cm2pm':'As,y média (cm²/m)',
    'laje_rho_l':'Taxa média introduzida diretamente',
    'edge_perp_interior':'Excentricidade perpendicular interior (bordo)',
    'corner_interior':'Excentricidades interiores (canto)',
    'simplified_applicable':'Condições de beta simplificado confirmadas',
    'opening_sectors':'Setores manuais adicionais: início/fim (graus)',
    'openings':'Aberturas definidas por posição e dimensões (m)',
    'reinforcement_diameter_mm':'Diâmetro dos ramos (mm)',
    'reinforcement_sr_m':'sr indicado (m; vazio = automático)',
    'reinforcement_s0_m':'s0 indicado (m; vazio = automático)',
    'cover_mm':'Recobrimento nominal (mm)','aggregate_mm':'Dimensão máxima do agregado (mm)','anchorage_confirmed':'Amarração confirmada pelo projetista',
    'footing_shape':'Forma da sapata','footing_bx':'Bx sapata (m)','footing_by':'By sapata (m)',
    'footing_diameter':'Diâmetro da sapata (m)'}


def fmt(v,nd=3):
    if v is None:return 'Não calculado'
    if isinstance(v,bool):return 'Sim' if v else 'Não'
    if isinstance(v,(int,float)):return f'{v:.{nd}f}' if math.isfinite(v) else 'Inválido'
    if isinstance(v,(list,dict,tuple)):return json.dumps(v,ensure_ascii=False)
    return str(v) or '-'


def sections(r):
    v=r['values'];p=r['inputs']
    results=[{'title':'Identificação e conclusão','lines':[
        f"Projeto: {p.get('project') or '-'} | Apoio: {p.get('support') or '-'} | Combinação: {p.get('combination') or '-'}",
        f"Versão: {r['version']} | Execução: {r['calculated_at']}",r['standard'],
        f"Estado: {r['badge']} ({r['status']})",r['conclusion'],f"Identificador: {r['input_hash'][:16]}"]},
        {'title':'Dados de cálculo','lines':[
            (f"Pilar {p['pilar_tipo']}, circular: D = {fmt(p['pilar_c1'])} m." if p['pilar_forma']=='circular' else f"Pilar {p['pilar_tipo']}, retangular: c1 = {fmt(p['pilar_c1'])} m; c2 = {fmt(p['pilar_c2'])} m."),
            f"d = {fmt(v['d'])} m; dx = {fmt(v['dx'])} m; dy = {fmt(v['dy'])} m.",
            f"fck = {fmt(p['betão_fck'])} MPa; fyk = {fmt(p['aço_fyk'])} MPa; fywk = {fmt(p['aço_fywk'])} MPa.",
            f"gamma_C = {fmt(p['gamma_C'])}; gamma_S = {fmt(p['gamma_S'])}.",
            f"VEd = {fmt(float(p['V_Ed'])/1000)} kN; MEdx = {fmt(float(p['M_Edx'])/1000)} kN.m; MEdy = {fmt(float(p['M_Edy'])/1000)} kN.m.",
            f"As,x = {fmt(p['laje_As_lx_cm2pm'])} cm²/m; As,y = {fmt(p['laje_As_ly_cm2pm'])} cm²/m.",
            f"rho_l = {fmt(v['rho_l']*100)} %; sigma_cp = {fmt(p['sigma_cp'])} MPa.",
            f"Aberturas: setores ineficazes combinados (graus) = {fmt(r['geometry'].get('opening_sectors',p['opening_sectors'] or []))}."]},
        {'title':'Geometria e coeficientes','lines':[
            f"u0 = {fmt(v['u0'])} m; u1 geométrico = {fmt(v['u1'])} m; u1 efetivo = {fmt(v['u1_eff'])} m.",
            (f"u1* = {fmt(v['u1_star'])} m; beta = {fmt(v['beta'])}." if v['u1_star'] is not None else f"u1* não utilizado neste caso; beta = {fmt(v['beta'])}."),
            r['beta_terms'].get('method','Método não avaliado.'),
            r.get('u0_terms',{}).get('method',''),
            f"k (efeito de escala) = {fmt(v['k_val'])}; C_Rd,c = {fmt(v['C_Rd_c'])}; v_min = {fmt(v['v_min'])} MPa.",
            f"vRd,c = {fmt(v['v_Rd_c'])} MPa; vRd,max = {fmt(v['v_Rd_max'])} MPa; kmax = {fmt(v['kmax'])}."]}]
    checks=[]
    if r.get('longitudinal_trace'):
        from .longitudinal import trace_lines as longitudinal_lines
        results.insert(2,{'title':'Armadura longitudinal e altura útil automática','lines':longitudinal_lines(r['longitudinal_trace'])})
    if r.get('load_trace'):
        from .connections import trace_lines
        results.append({'title':'Origem e equilíbrio dos esforços da ligação','lines':trace_lines(r['load_trace'])})
    terms=r['beta_terms']
    if terms.get('steps'):
        lines=[terms['method'], 'Referência: '+terms['reference']]
        lines.extend(terms.get('applicability',[]))
        for step in terms['steps']:
            lines.append(f"{step['symbol']}: {step['expression']}. Substituição: {step['substitution']} = {fmt(step['value'],6)} {step['unit']}. [{step['reference']}]")
        lines.append('O beta acima é o utilizado na verificação de u1 e da face. Na pesquisa do contorno exterior conserva-se pelo menos este valor; uma majoração adicional devida ao contorno exterior é uma opção conservadora do programa.')
        results.append({'title':'Determinação do coeficiente beta','lines':lines})
    if r.get('opening_records'):
        lines=['Posições X/Y medidas ao centro do pilar. Distâncias mínimas face a face; setores automáticos unidos aos setores manuais. Referência: 6.4.2(3), Figura 6.14.']
        for o in r['opening_records']:
            size=f"D = {fmt(o['diameter'])} m" if o['shape']=='circular' else f"dimensões X/Y = {fmt(o['width'])} / {fmt(o['height'])} m"
            lines.append(f"{o['id']}: centro X/Y = {fmt(o['x'])} / {fmt(o['y'])} m; {size}. Distância livre = {fmt(o['distance_m'])} m; limite 6d = {fmt(o['limit_6d_m'])} m.")
            lines.append(o['method'])
            if o['active']:
                lines.append(f"Setor considerado: {fmt(o['sector_deg'][0],6)} a {fmt(o['sector_deg'][1],6)} graus, no sentido positivo de X para Y.")
                if o['equivalent_outline']:lines.append(f"Envolvente radial: l1 = {fmt(o['l1_m'])} m; l2 = {fmt(o['l2_m'])} m; largura equivalente = sqrt(l1*l2) = {fmt(o['equivalent_width_m'])} m.")
        results.append({'title':'Aberturas - definição geométrica e setores','lines':lines})
    if p.get('edge_distance_m',0)>0:
        lines=[f"Afastamento livre g = {fmt(p['edge_distance_m'])} m. c1 paralelo e c2 perpendicular ao bordo. A distância é medida à face do pilar."]
        for c in r.get('contour_comparison',[]):
            name='Contorno fechado' if c['kind']=='fechado' else 'Contorno aberto ao bordo'
            state=('adotado como mínimo' if c['selected'] else 'admissível') if c['admissible'] else 'não admissível: toca ou ultrapassa o bordo'
            lines.append(f"{name}: comprimento geométrico = {fmt(c['geometric_length'])} m; efetivo = {fmt(c['effective_length'])} m; {state}.")
            if c['admissible']:
                detail=next((x for x in r['beta_terms'].get('candidates',[]) if x['kind']==c['kind']),None)
                if detail is not None:
                    lines.append(f"beta = {fmt(c.get('beta'))}; vEd = {fmt(c.get('vEd'))} MPa; centro do perímetro (x;y) = ({fmt(detail['cx'])}; {fmt(detail['cy'])}) m; Wx = {fmt(detail['Wx'])} m²; Wy = {fmt(detail['Wy'])} m².")
                else:lines.append('Beta e tensão neste contorno não calculados; consulte o estado e as notas da execução.')
        if r['beta_terms'].get('source') in ('manual','simplificado'):
            lines.append('O mesmo beta declarado é aplicado aos contornos e à face. O contorno mínimo condiciona a tensão. CG e W descrevem a geometria; não se subtrai VEd*CG aos momentos.')
        results.append({'title':'Afastamento ao bordo e comparação de contornos','lines':lines})
    for c in r['checks']:
        checks.append(f"{c['name']}: vEd = {fmt(c['demand'])} MPa; resistência = {fmt(c['resistance'])} MPa; utilização = {fmt(c['utilization'])}; {'VERIFICA' if c['passed'] else 'NÃO VERIFICA'}. [{c['reference']}]")
    results.append({'title':'Verificações realizadas','lines':checks or ['Nenhuma verificação concluída.']})
    if p['is_sapata']:
        points=r['footing_scan'];peak=max(points,key=lambda x:x['ratio']) if points else None
        results.append({'title':'Verificação da sapata','lines':[
            (f"Sapata circular: D = {fmt(p['footing_diameter'])} m." if p['footing_shape']=='circular' else f"Sapata retangular: Bx = {fmt(p['footing_bx'])} m; By = {fmt(p['footing_by'])} m."),
            f"Pressão líquida de equilíbrio = {fmt(v.get('soil_pressure_used_kpa'))} kPa.",
            f"Número de secções avaliadas = {len(points)}; distância condicionante a = {fmt(v['a_governing'],6)} m.",
            ('Na secção condicionante: u = '+fmt(peak['u'])+' m; área de reação = '+fmt(peak['area'])+' m²; VEd,red = '+fmt(peak['V_red']/1000)+' kN.') if peak else 'Secção condicionante não determinada.',
            'A listagem completa dos perímetros integra o XLSX e o JSON de resultados.']})
    if r['reinforcement_rows']:
        lines=[f"Estribos verticais: diâmetro {fmt(p['reinforcement_diameter_mm'],0)} mm. Amarração confirmada: {fmt(p['anchorage_confirmed'])}.",
               f"fywd,ef = {fmt(v['f_ywd_ef'])} MPa; Asw/sr de cálculo = {fmt(v['Asw_sr_calc']*1e4)} cm²/m.",
               f"Asw/sr de referência para o mínimo em u1 = {fmt(v['Asw_sr_min_u1_reference']*1e4)} cm²/m. Não é o mínimo a impor a todas as fiadas.",
               f"Asw/sr requerido por resistência (6.52) = {fmt(v['Asw_sr_req']*1e4)} cm²/m. O mínimo (9.11) é verificado por ramo com sr e st efetivos de cada fiada.",
               f"u_out,ef = {fmt(v['u_out_ef'])} m; distância à face = {fmt(v['r_out_face'])} m.",
               f"Distância mínima da última fiada à face = {fmt(v['dist_zona_armar'])} m."]
        for row in r['reinforcement_rows']:
            lines.append(f"Fiada {row['row']}: r = {fmt(row['r'])} m; {row['n_legs']} ramos de {fmt(row['phi_mm'],0)} mm; Asw total = {fmt(row['Asw_m2']*1e4)} cm²; sr = {fmt(row['sr'])} m; st máximo = {fmt(row['st'])} m (limite {fmt(row['st_max'])} m).")
            lines.append(f"Mínimo local (9.11): área por ramo = {fmt(row['Asw_branch_provided_m2']*1e4)} cm²; mínimo por ramo = {fmt(row['Asw_min_branch_m2']*1e4)} cm²; utilização = {fmt(row['minimum_branch_utilization'])}. Asw requerido nesta fiada = {fmt(row['Asw_required_m2']*1e4)} cm² (resistência e mínimo local).")
            for family in row.get('families',[]):
                name='Contorno aberto' if family['kind']=='aberto_bordo' else 'Contorno fechado'
                lines.append(f"{name} da fiada {row['row']}: {family['n_legs']} ramos contribuem; Asw = {fmt(family['Asw_m2']*1e4)} cm² (requerido {fmt(family['Asw_required_m2']*1e4)} cm²); st = {fmt(family['st'])} m; {'VERIFICA' if family['passed'] else 'NÃO VERIFICA'}.")
        if any(row.get('families') for row in r['reinforcement_rows']):
            lines.append('Asw total é o inventário físico da fiada. Os ramos comuns contribuem para ambos os contornos alternativos, mas não se somam duas vezes no inventário. A correspondência entre ramos e contornos está no XLSX e no JSON.')
        results.append({'title':'Solução de armadura e pormenorização','lines':lines})
    elif r['status']=='EDGE_DETAIL_PENDING':
        lines=[f"fywd,ef = {fmt(v['f_ywd_ef'])} MPa; Asw/sr de cálculo = {fmt(v['Asw_sr_calc']*1e4)} cm²/m.",
               f"Asw/sr mínimo de referência = {fmt(v['Asw_sr_min']*1e4)} cm²/m; requerido por resistência = {fmt(v['Asw_sr_req']*1e4)} cm²/m. Confirmar o mínimo local por ramo.",
               f"u_out,ef = {fmt(v['u_out_ef'])} m; distância de referência à face = {fmt(v['r_out_face'])} m.",
               f"Distância mínima de referência da última fiada à face = {fmt(v['dist_zona_armar'])} m.",
               'Não foi definida nem aprovada uma distribuição de ramos. O pormenor deve verificar a armadura que atravessa cada contorno, os espaçamentos, os extremos junto ao bordo e a amarração. A confirmação genérica de amarração não elimina esta pendência.']
        for c in r.get('reinforcement_demands',[]):
            lines.append(f"{c['contour']}: vEd = {fmt(c['vEd'])} MPa; requer armadura = {fmt(c['requires_steel'])}; Asw/sr requerido = {fmt(c['Asw_sr_req_m2pm']*1e4)} cm²/m.")
        results.append({'title':'Armadura necessária - pormenor específico pendente','lines':lines})
    elif r['status']=='PASS_WITHOUT':
        results.append({'title':'Armadura de punçoamento','lines':['Não é necessária armadura específica nos mecanismos e no domínio avaliados. Manter as armaduras longitudinais e de integridade previstas no projeto.']})
    else:
        results.append({'title':'Armadura de punçoamento','lines':['Não foi obtida uma solução de armadura aprovada para esta execução. Valores não calculados não representam armadura nula.']})
    if r.get('detail_checks'):
        lines=[]
        for c in r['detail_checks']:
            detail=(f"{fmt(c['value'])} {c['rule']} {fmt(c['limit'])} {c['unit']}" if c['id'] not in ('paths','families','radii') else 'conferência das posições propostas')
            lines.append(f"{c['name']}: {detail}; {'VERIFICA' if c['passed'] else 'NÃO VERIFICA'}.")
        results.append({'title':'Conferência da distribuição proposta','lines':lines})
    results.append({'title':'Hipóteses, referências e notas','lines':r['notes']+[
        'Pormenorização de estribos: 9.4.3 e Figura 9.10. Limite resistente: A1:2019, 6.4.5(1). Escora: AC:2012, correção 95.',
        REPOSITORY]})
    return results


def text_report(r):
    parts=['PunchingShearEC2 - Memória de verificação de punçoamento']
    for sec in sections(r):parts.extend(['',sec['title']]+sec['lines'])
    return '\n'.join(parts)+'\n'


def export_json(r,path):Path(path).write_text(json.dumps(r,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def export_txt(r,path):Path(path).write_text(text_report(r),encoding='utf-8')


def export_xlsx(r,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();wb.remove(wb.active);wb.properties.creator=REPOSITORY
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    def sheet(name,headers,rows):
        ws=wb.create_sheet(name);ws.append(headers)
        for row in rows:
            # User metadata is text, never executable spreadsheet formulas.
            ws.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
        for row in ws:
            for cell in row:
                cell.font=Font(name='Courier New',size=10)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,float):cell.number_format='0.000'
        for cell in ws[1]:
            cell.font=Font(name='Courier New',size=10,bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='173C4B')
        for i in range(1,ws.max_column+1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(i)].width=50 if i==1 else 24
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        return ws
    sheet('Resumo',['Campo','Valor'],[
        ['Estado',r['status']],['Conclusão',r['conclusion']],['Versão',r['version']],['Norma',r['standard']],
        ['Execução',r['calculated_at']],['Identificador',r['input_hash']]])
    sheet('Entradas',['Campo','Valor'],[[INPUT_LABELS.get(k,k),fmt(v) if isinstance(v,(list,dict,tuple)) else v] for k,v in r['inputs'].items()])
    if r.get('longitudinal_trace'):
        from .longitudinal import trace_lines as longitudinal_lines
        sheet('ArmaduraLongitudinal',['Armadura longitudinal e altura útil'],[[line] for line in longitudinal_lines(r['longitudinal_trace'])]).column_dimensions['A'].width=110
    if r.get('load_trace'):
        from .connections import trace_lines
        ws_origin=sheet('OrigemEsforcos',['Origem e equilíbrio dos esforços'],[[line] for line in trace_lines(r['load_trace'])])
        ws_origin.column_dimensions['A'].width=110
    ws=sheet('Verificacoes',['Verificação','vEd (MPa)','Resistência (MPa)','Utilização motor','Utilização XLSX','Estado','Referência'],
          [[c['name'],c['demand'],c['resistance'],c['utilization'],None,'VERIFICA' if c['passed'] else 'NÃO VERIFICA',c['reference']] for c in r['checks']])
    for i in range(2,ws.max_row+1):ws.cell(i,5,f'=IFERROR(B{i}/C{i},"")')
    sheet('Parametros',['Variável','Valor'],list(r['values'].items())+[['beta: '+k,fmt(v) if isinstance(v,(list,dict)) else v] for k,v in r['beta_terms'].items() if k!='candidates'])
    sheet('Beta',['Grandeza','Expressão','Substituição numérica','Valor','Unidade','Referência'],
          [[s['symbol'],s['expression'],s['substitution'],s['value'],s['unit'],s['reference']] for s in r['beta_terms'].get('steps',[])])
    if r.get('opening_records'):
        sheet('Aberturas',['Abertura','Forma','Centro X (m)','Centro Y (m)','Dimensão X (m)','Dimensão Y (m)','D (m)','Distância livre (m)','6d (m)','Setor considerado','Ângulo inicial (graus)','Ângulo final (graus)','l1 radial (m)','l2 transversal (m)','Largura equivalente (m)','Método','Referência'],
            [[o['id'],'Circular' if o['shape']=='circular' else 'Retangular',o['x'],o['y'],o.get('width'),o.get('height'),o.get('diameter'),o['distance_m'],o['limit_6d_m'],o['active'],
              *(o['sector_deg'] or [None,None]),o['l1_m'],o['l2_m'],o['equivalent_width_m'],o['method'],o['reference']] for o in r['opening_records']])
    if r.get('contour_comparison'):
        details={c['kind']:c for c in r['beta_terms'].get('candidates',[])}
        sheet('Contornos',['Contorno','u geom. (m)','u efetivo (m)','Admissível','Mínimo adotado','Beta adotado','vEd (MPa)','Condiciona tensão','CG X (m)','CG Y (m)','Wx (m²)','Wy (m²)'],
              [[c['kind'],c['geometric_length'],c['effective_length'],c['admissible'],c['selected'],c.get('beta'),c.get('vEd'),c.get('stress_governing'),
                *[details.get(c['kind'],{}).get(k) for k in ('cx','cy','Wx','Wy')]] for c in r['contour_comparison']])
    if r.get('reinforcement_demands'):
        sheet('ArmaduraNecessaria',['Contorno','u (m)','Beta adotado','vEd (MPa)','Requer armadura','Asw/sr cálculo (cm²/m)','Asw/sr mínimo de referência neste contorno (cm²/m)','Asw/sr requerido por resistência (cm²/m)','Asw/sr fornecido no contorno (cm²/m)','Resistência (MPa)','Utilização'],
              [[c['contour'],c['u'],c['beta_adopted'],c['vEd'],c['requires_steel'],c['Asw_sr_calc_m2pm']*1e4,c['Asw_sr_min_m2pm']*1e4,c['Asw_sr_req_m2pm']*1e4,
                c['Asw_sr_provided_m2pm']*1e4 if c.get('Asw_sr_provided_m2pm') is not None else None,c.get('resistance_MPa'),c.get('utilization')] for c in r['reinforcement_demands']])
    sheet('Fiadas',['Fiada','r face (m)','N ramos únicos','Diâmetro (mm)','Asw total (cm²)','Asw requerido por contorno (cm²)','sr (m)','st máximo (m)','st limite (m)','Verifica localmente','Área por ramo (cm²)','Mínimo por ramo (cm²)','Utilização do mínimo (9.11)','Asw por resistência (cm²)','Asw mínimo local de referência (cm²)'],
          [[x['row'],x['r'],x['n_legs'],x['phi_mm'],x['Asw_m2']*1e4,x['Asw_required_m2']*1e4,x['sr'],x['st'],x['st_max'],x['passed'],x['Asw_branch_provided_m2']*1e4,x['Asw_min_branch_m2']*1e4,x['minimum_branch_utilization'],x['Asw_resistance_required_m2']*1e4,x['Asw_min_local_m2']*1e4] for x in r['reinforcement_rows']])
    if any(row.get('families') for row in r['reinforcement_rows']):
        sheet('FiadasContornos',['Fiada','Contorno','Ramos que contribuem','N ramos','Asw (cm²)','Asw requerido (cm²)','st (m)','st limite (m)','Asw mínimo por ramo (cm²)','Verifica'],
              [[row['row'],f['kind'],', '.join(map(str,f['leg_indices'])),f['n_legs'],f['Asw_m2']*1e4,f['Asw_required_m2']*1e4,f['st'],f['st_max'],f['Asw_min_branch_m2']*1e4,f['passed']]
               for row in r['reinforcement_rows'] for f in row.get('families',[])])
    if r.get('detail_checks'):
        sheet('Pormenorizacao',['Condição','Valor','Regra','Limite','Unidade','Verifica'],
              [[c['name'],c['value'],c['rule'],c['limit'],c['unit'],c['passed']] for c in r['detail_checks']])
    sheet('Ramos',['Fiada','Ramo','X (m)','Y (m)','Diâmetro (mm)'],
          [[row['row'],i+1,x,y,row['phi_mm']] for row in r['reinforcement_rows'] for i,(x,y) in enumerate(row['coordinates'])])
    if r['footing_scan']:
        sheet('Sapatas',['a (m)','u (m)','Área reação (m²)','VEd red (N)','vEd (MPa)','vRd (MPa)','Utilização'],
              [[p[k] for k in ['a','u','area','V_red','vEd','vRd','ratio']] for p in r['footing_scan']])
    sheet('Notas',['Notas'],[[n] for n in r['notes']])
    wb.save(path)


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os
    system_fonts=Path(os.environ.get('SystemRoot','/'))/'Fonts'
    pairs=[(system_fonts/'cour.ttf',system_fonts/'courbd.ttf'),
           (Path('/Library/Fonts/Courier New.ttf'),Path('/Library/Fonts/Courier New Bold.ttf')),
           (Path('/usr/share/fonts/truetype/liberation2/LiberationMono-Regular.ttf'),Path('/usr/share/fonts/truetype/liberation2/LiberationMono-Bold.ttf')),
           (Path('/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf'))]
    for regular,bold in pairs:
        if regular.is_file() and bold.is_file():
            if 'EC2Mono' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('EC2Mono',str(regular)))
                pdfmetrics.registerFont(TTFont('EC2MonoBold',str(bold)))
                pdfmetrics.registerFontFamily('EC2Mono',normal='EC2Mono',bold='EC2MonoBold')
            return 'EC2Mono','EC2MonoBold'
    return 'Courier','Courier-Bold'


def drawing(r,width=482,height=290):
    from reportlab.graphics.shapes import Drawing,PolyLine,Rect,Circle,String,Line
    from reportlab.lib import colors
    d=Drawing(width,height);g=r['geometry'];c=g.get('column')
    if not c:return d
    layers=[('u0','#6b46a0'),('u1','#c53030'),('u1_star','#b7770a'),('u_out','#19835b'),('governing','#176eaa')]
    points=[point for key,_ in layers for path in g.get(key,[]) for point in path]
    points.extend([(-c['c1']/2,-c['c2']/2),(c['c1']/2,c['c2']/2)])
    for opening in g.get('openings',[]):points.extend(opening['outline']+opening['equivalent_outline'])
    if c['position']=='bordo':points.append((0,-c['c2']/2-c.get('edge_distance_m',0)))
    if g.get('footing'):
        f=g['footing'];points.extend([(-f['bx']/2,-f['by']/2),(f['bx']/2,f['by']/2)])
    minx=min(x for x,y in points);maxx=max(x for x,y in points)
    miny=min(y for x,y in points);maxy=max(y for x,y in points)
    scale=min((width-90)/max(maxx-minx,.1),(height-60)/max(maxy-miny,.1))
    ox=width/2-scale*(minx+maxx)/2;oy=(height+20)/2-scale*(miny+maxy)/2
    def pt(x,y):return ox+x*scale,oy+y*scale
    if g.get('footing'):
        f=g['footing']
        if f['shape']=='circular':d.add(Circle(ox,oy,f['bx']/2*scale,strokeColor=colors.lightgrey,fillColor=None))
        else:d.add(Rect(*pt(-f['bx']/2,-f['by']/2),f['bx']*scale,f['by']*scale,strokeColor=colors.lightgrey,fillColor=None))
    if c['position'] in ('bordo','canto'):
        gap=c.get('edge_distance_m',0)
        yy=pt(0,-c['c2']/2-gap)[1];d.add(Line(15,yy,width-15,yy,strokeColor=colors.grey,strokeWidth=1.4))
        if gap>0:
            xx,y=pt(0,-c['c2']/2)
            col=colors.HexColor('#176eaa');d.add(Line(xx,y,xx,yy,strokeColor=col))
            for level in (y,yy):d.add(Line(xx-3,level,xx+3,level,strokeColor=col))
            d.add(String(xx+6,(y+yy)/2,f'g = {gap:.3f} m',fontName=_pdf_fonts()[0],fontSize=8,fillColor=col))
    if c['position']=='canto':
        xx=pt(-c['c1']/2,0)[0];d.add(Line(xx,30,xx,height-10,strokeColor=colors.grey,strokeWidth=1.4))
    if c['shape']=='circular':d.add(Circle(ox,oy,c['c1']/2*scale,fillColor=colors.HexColor('#e3ebf0'),strokeColor=colors.HexColor('#173c4b')))
    else:d.add(Rect(*pt(-c['c1']/2,-c['c2']/2),c['c1']*scale,c['c2']*scale,fillColor=colors.HexColor('#e3ebf0'),strokeColor=colors.HexColor('#173c4b')))
    for key,color in layers:
        for path in g.get(key,[]):
            d.add(PolyLine([z for x,y in path for z in pt(x,y)],strokeColor=colors.HexColor(color),strokeWidth=1.3,fillColor=None))
    from .openings import dimensions
    for o in g.get('openings',[]):
        color=colors.HexColor('#a35d05') if o['active'] else colors.grey
        for end in o['tangent_points']:d.add(Line(*pt(0,0),*pt(*end),strokeColor=color,strokeDashArray=[3,3],strokeWidth=.6))
        if o['equivalent_outline']:d.add(PolyLine([z for p in o['equivalent_outline'] for z in pt(*p)],strokeColor=color,strokeDashArray=[3,3],fillColor=None))
        w,h=dimensions(o)
        if o['shape']=='circular':d.add(Circle(*pt(o['x'],o['y']),w/2*scale,strokeColor=color,fillColor=colors.HexColor('#fff3df')))
        else:d.add(Rect(*pt(o['x']-w/2,o['y']-h/2),w*scale,h*scale,strokeColor=color,fillColor=colors.HexColor('#fff3df')))
        x,y=pt(o['x'],o['y']);d.add(String(x,y,o['id'] if len(o['id'])<=12 else o['id'][:9]+'...',fontName=_pdf_fonts()[0],fontSize=8,textAnchor='middle',fillColor=color))
    for row in r['reinforcement_rows']:
        for x,y in row['coordinates']:d.add(Circle(*pt(x,y),1.7,fillColor=colors.HexColor('#173c4b'),strokeColor=None))
    font,_=_pdf_fonts()
    d.add(String(12,12,'u0 roxo | u1 vermelho | u1* ocre | u_out verde | secção crítica azul',fontName=font,fontSize=7))
    d.add(String(width-63,height-15,'+X →',fontName=font,fontSize=8))
    d.add(String(width-63,height-28,'+Y ↑',fontName=font,fontSize=8))
    return d


def export_pdf(r,path):
    from html import escape
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    font,bold=_pdf_fonts();body=ParagraphStyle('body',fontName=font,fontSize=10,leading=15,spaceAfter=6)
    heading=ParagraphStyle('heading',fontName=bold,fontSize=12,leading=18,spaceBefore=15,spaceAfter=12)
    small=ParagraphStyle('small',parent=body,fontSize=8,leading=10,spaceAfter=0)
    title=ParagraphStyle('title',parent=heading,fontSize=15,leading=20,spaceAfter=16)
    def par(text,style=body):return Paragraph(escape(str(text)),style)
    doc=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=52,rightMargin=52,topMargin=42,bottomMargin=60,
                         title='PunchingShearEC2 - Memória de cálculo',author=REPOSITORY)
    story=[par('Verificação de punçoamento',title)]
    for index,sec in enumerate(sections(r)):
        story.append(KeepTogether([par(sec['title'],heading),par(sec['lines'][0])]))
        for line in sec['lines'][1:]:story.append(par(line))
        if index==2 and r['geometry']:
            story.append(Spacer(1,8));story.append(drawing(r,width=A4[0]-104))
    rows=[[par('Entrada',small),par('Valor',small)]]
    p=r['inputs']
    ignored={'u1_ineffective','openings'}
    if not p['is_sapata']:ignored.update({'footing_shape','footing_bx','footing_by','footing_diameter','sigma_gd_kpa'})
    if p['pilar_tipo']!='bordo':ignored.add('edge_perp_interior')
    if p['pilar_tipo']!='canto':ignored.add('corner_interior')
    if p['pilar_forma']=='circular':ignored.add('pilar_c2')
    for k,v in p.items():
        if v is not None and k not in ignored:rows.append([par(INPUT_LABELS.get(k,k),small),par(fmt(v),small)])
    table=Table(rows,colWidths=[295,A4[0]-104-295],repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e3ebf0')),
                              ('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.6,colors.grey),
                              ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
                              ('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),2)]))
    story.append(KeepTogether([par('Registo das entradas do caso',heading),table]))
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont(font,8)
        canvas.setStrokeColor(colors.lightgrey);canvas.line(52,46,A4[0]-52,46)
        canvas.drawCentredString(A4[0]/2,34,f"PunchingShearEC2 {r['version']} | {r['calculated_at']} | {doc.page}")
        canvas.setFont(font,8);canvas.drawCentredString(A4[0]/2,22,REPOSITORY)
        canvas.linkURL(REPOSITORY,(52,18,A4[0]-52,29),relative=0)
        canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
