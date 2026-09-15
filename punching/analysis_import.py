"""Modelo bar-table adapter. Reading records is separate from joint equilibrium.

Internal moments are local sectional forces, NOT actions on the joint. For 3D
bars: joint moment = +R*M at the end, -R*M at the beginning. Only vertical bars
and local transverse axes parallel to the calculation axes are supported here.
Nothing infers a member direction, missing upper column, ELU, or axis rotation.
"""
from copy import deepcopy
from dataclasses import asdict
import csv
import io
import json
import math
from pathlib import Path
import re

from .table_import import Table, read_table, token, numeric, identity, key, MAX_ROWS

ADAPTER = 'Modelo.bar-joint/1'
SIGN_REFERENCE = 'Perfil de barras 3D: FX positivo em compressão; momentos locais convertidos em ações sobre o nó (+ no fim, - no início). Confira estes sinais na exportação de origem.'
COMPOUND = ('member/node/case', 'bar/node/case', 'barra/no/caso', 'barra/no/combinacao')
ALIASES = {'name': ('name', 'nome'), 'story': ('story', 'piso'),
           'node': ('node', 'no'), 'x': ('x',), 'y': ('y',), 'z': ('z',)}
UNITS = {'fx': {'kn': 1., 'n': .001},
         'my': {'knm': 1., 'nm': .001, 'nmm': .000001, 'kncm': .01},
         'mz': {'knm': 1., 'nm': .001, 'nmm': .000001, 'kncm': .01},
         'hy': {'m': 1., 'cm': .01, 'mm': .001},
         'hz': {'m': 1., 'cm': .01, 'mm': .001},
         'ax': {'m2': 1., 'cm2': .0001, 'mm2': .000001},
         'iy': {'m4': 1., 'cm4': 1e-8, 'mm4': 1e-12},
         'iz': {'m4': 1., 'cm4': 1e-8, 'mm4': 1e-12}}
AXES = {'+X': (1, 0), '-X': (-1, 0), '+Y': (0, 1), '-Y': (0, -1)}


def modal_combination_issue(combo):
    if not re.search(r'\b(CQC|SRSS)\b', combo, flags=re.IGNORECASE): return ''
    return ('Combinação CQC/SRSS: os resultados modais já combinados não comprovam concomitância. '
            'VEd e MEd não são obtidos por diferença/soma direta destes valores. '
            'Obtenha a resultante da ligação com tratamento modal adequado e importe-a como fonte fundamentada.')


def header(value):
    text = token(value).replace('²', '2').replace('⁴', '4')
    text = re.sub(r'[\s.^·*]', '', text)
    match = re.fullmatch(r'([^\[\(]+)(?:[\[(]([^\]\)]+)[\])])?', text)
    return match.groups() if match else (text, None)


def profile(headers):
    names = {header(h)[0] for h in headers}
    if any(x in names for x in COMPOUND) and {'fx', 'my', 'mz'} <= names: return 'analysis'
    if names.intersection(ALIASES['node']) and {'x', 'y', 'z'} <= names: return 'analysis_nodes'
    return None


def split_key(value):
    match = re.fullmatch(r'\s*([0-9]+)\s*/\s*([0-9]+)\s*/\s*(.+?)\s*', str(value))
    if not match: raise ValueError('Member/Node/Case: esperado barra/nó/caso, por exemplo 2/7460/101 (C).')
    member, node, case = match.groups()
    if int(member) <= 0 or int(node) <= 0: raise ValueError('Barra e nó devem ter números positivos.')
    case = identity(' '.join(case.split()), 'Caso / combinação')
    return member, node, case


def column_map(table):
    result = {}
    for i, h in enumerate(table.headers):
        name, unit = header(h)
        name = 'compound' if name in COMPOUND else next((k for k, v in ALIASES.items() if name in v), name)
        if name in result: raise ValueError(f'Coluna Modelo ambígua: {h}.')
        result[name] = (i, unit)
    return result


def _number(value, label, decimal):
    # Modelo can use a space/nonbreaking space as the thousands separator.
    if isinstance(value, str):
        value = value.strip().replace('\u00a0', ' ').replace('\u202f', ' ').replace('\u2212', '-')
        if ' ' in value:
            if not re.fullmatch(r'[+-]?\d{1,3}(?: \d{3})+(?:' + re.escape(decimal) + r'\d+)?', value):
                raise ValueError(f'{label}: agrupamento de milhares inválido.')
            value = value.replace(' ', '')
    return numeric(value, label, decimal)


def detect_decimal(table, mapping):
    separators = set()
    for _, row in table.rows:
        for name in ('fx', 'my', 'mz', 'hy', 'hz', 'x', 'y', 'z'):
            if name not in mapping: continue
            value = row[mapping[name][0]]
            if isinstance(value, str):
                if ',' in value: separators.add(',')
                if '.' in value: separators.add('.')
    if len(separators) > 1: raise ValueError('A tabela de barras mistura ponto e vírgula decimal. Exporte valores com uma convenção única, sem milhares com ponto/vírgula.')
    return next(iter(separators), '.')


def automatic_layout(path, sheet=None):
    """Recognize known headers in the first 50 rows; generic layouts stay manual."""
    path = Path(path)
    if path.stat().st_size > 32*1024*1024: raise ValueError('A tabela excede 32 MB.')
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = wb[sheet or wb.sheetnames[0]]; ws.reset_dimensions()
            for i, row in enumerate(ws.iter_rows(max_row=50, max_col=256, values_only=True), 1):
                if profile(row): return {'header_row': i}
        finally: wb.close()
    else:
        payload = path.read_bytes()
        encoding = 'utf-16' if payload.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        try: contents = payload.decode(encoding)
        except UnicodeError:
            encoding = 'cp1252'; contents = payload.decode(encoding)
        for delimiter in ('\t', ';', ','):
            for i, row in enumerate(csv.reader(io.StringIO(contents), delimiter=delimiter), 1):
                if i > 50: break
                if profile(row): return {'header_row': i, 'delimiter': delimiter, 'encoding': encoding}
    return {}


def archive(table):
    data = asdict(table); data['formulas'] = [list(x) for x in sorted(table.formulas)]
    return json.loads(json.dumps(data, ensure_ascii=False, allow_nan=False))


def restore(data):
    table = Table(**deepcopy(data)); table.formulas = {tuple(x) for x in table.formulas}
    if len(table.rows) > MAX_ROWS or not 1 <= len(table.headers) <= 256: raise ValueError('Tabela de barras demasiado grande.')
    if any(len(raw) != len(table.headers) for _, raw in table.rows): raise ValueError('Linha Modelo incompleta.')
    return table


def _source(table, row, raw):
    return {'file': table.file, 'sha256': table.sha256, 'sheet': table.sheet, 'row': row,
            'raw': dict(zip(table.headers, raw))}


def _section(values):
    if not all(k in values for k in ('hy', 'hz', 'ax', 'iy', 'iz')):
        return None, 'Secção por confirmar: exporte HY, HZ, AX, IY e IZ com unidades.'
    hy, hz, area, iy, iz = (values[k] for k in ('hy', 'hz', 'ax', 'iy', 'iz'))
    if not (.01 <= hy <= 20 and .01 <= hz <= 20) or min(area, iy, iz) <= 0:
        return None, 'Dimensões ou propriedades da secção fora do domínio.'
    def close(a, b): return math.isclose(a, b, rel_tol=.005, abs_tol=1e-12)
    if close(area, hy*hz) and close(iy, hy*hz**3/12) and close(iz, hz*hy**3/12):
        return {'shape': 'retangular', 'hy': hy, 'hz': hz}, ''
    if close(hy, hz) and close(area, math.pi*hy**2/4) and close(iy, math.pi*hy**4/64) and close(iz, iy):
        return {'shape': 'circular', 'hy': hy, 'hz': hz}, ''
    return None, 'HY/HZ e propriedades AX/IY/IZ não confirmam uma secção maciça retangular ou circular (tolerância 0,5%).'


def parse_bars(table):
    if profile(table.headers) != 'analysis': raise ValueError('Tabela de esforços de barras Modelo não reconhecida.')
    mapping = column_map(table)
    for name in ('compound', 'story', 'fx', 'my', 'mz'):
        if name not in mapping: raise ValueError(f'Exporte também a coluna {name.upper()} na tabela de barras.')
    for name, units in UNITS.items():
        if name in mapping and mapping[name][1] not in units:
            raise ValueError(f'{table.headers[mapping[name][0]]}: unidade em falta ou não suportada. Inclua a unidade no cabeçalho.')
    decimal = detect_decimal(table, mapping); records = []; seen = set(); errors = []
    used = set(UNITS).intersection(mapping) | {'compound', 'story'} | ({'name'} if 'name' in mapping else set())
    for row, raw in table.rows:
        try:
            if any((row, mapping[k][0]) in table.formulas for k in used): raise ValueError('Fórmula/erro XLSX: exporte os valores calculados.')
            member, node, combo = split_key(raw[mapping['compound'][0]])
            rkey = key(member, node, combo)
            if rkey in seen: raise ValueError(f'Barra/nó/caso repetido: {member}/{node}/{combo}. Não se somam duplicados.')
            seen.add(rkey)
            values = {k: _number(raw[mapping[k][0]], k.upper(), decimal)*UNITS[k][mapping[k][1]] for k in set(UNITS).intersection(mapping)}
            section, issue = _section(values)
            supplied_name = raw[mapping['name'][0]] if 'name' in mapping else None
            missing_name = supplied_name is None or isinstance(supplied_name, str) and not supplied_name.strip()
            records.append({'member': member, 'node': node, 'combination': combo,
                'support': f'Barra {member}' if missing_name else identity(supplied_name, 'Name / pilar'),
                'story': identity(raw[mapping['story'][0]], 'Story / piso da barra'),
                'values': values, 'section': section, 'section_issue': issue, 'source': _source(table, row, raw)})
            if missing_name: records[-1]['name_resolution'] = {'method': 'member_id'}
        except (ValueError, TypeError) as exc: errors.append(f'Linha {row}: {exc}')
    if errors: raise ValueError('\n'.join(errors[:20]) + (f'\nMais {len(errors)-20} erros.' if len(errors)>20 else ''))
    if not records: raise ValueError('Tabela sem esforços.')
    _resolve_names(records)
    members(records)  # identity and section consistency across all rows
    return records


def _resolve_names(records):
    """Names are labels; bar/node/case remains the structural identity.

    Recover a blank label only from explicit names on the SAME member. Preserve
    the evidence row for compact case archives. Never fill from a neighbouring
    row, another connected member, a section size or a similar support name.
    """
    explicit = {}
    for r in records:
        if 'name_resolution' not in r:
            names = explicit.setdefault(r['member'], {})
            names.setdefault(r['support'], r['source'])
            if len(names) > 1: raise ValueError(f"Barra {r['member']}: nomes de pilar incompatíveis nas linhas de origem. Confira a identificação no modelo.")
    for r in records:
        if 'name_resolution' not in r: continue
        names = explicit.get(r['member'], {})
        if names:
            label, evidence = next(iter(names.items()))
            r['support'] = label
            r['name_resolution'] = {'method': 'member_name', 'evidence': deepcopy(evidence)}
        else:
            r['support'] = f"Barra {r['member']}"
            r['name_resolution'] = {'method': 'member_id'}


def name_note(record):
    resolution = record.get('name_resolution')
    if not resolution: return ''
    if resolution['method'] == 'member_id':
        return f"Name não preenchido na origem. Identificação de consulta: «Barra {record['member']}», obtida do número da barra. Pode definir o nome da ligação ao prepará-la."
    source = resolution['evidence']
    return f"Name não preenchido nesta linha. Nome «{record['support']}» recuperado de outra linha da mesma barra {record['member']}: {source['file']}, folha {source['sheet']}, linha {source['row']}."


def name_summary(records):
    missing = [r for r in records if r.get('name_resolution')]
    if not missing: return []
    fallback = {r['member'] for r in missing if r['name_resolution']['method'] == 'member_id'}
    recovered = {r['member'] for r in missing if r['name_resolution']['method'] == 'member_name'}
    lines = [f"Name não preenchido em {len(missing)} registos, correspondentes a {len({r['member'] for r in missing})} barras. Todos os registos foram conservados."]
    if fallback:
        labels = ', '.join('Barra '+m for m in sorted(fallback, key=lambda m:(int(m), m))[:20])
        lines.append(f"{len(fallback)} barras identificadas pelo número: {labels}" + ('…' if len(fallback)>20 else '.') + ' O nome da ligação pode ser definido ao prepará-la.')
    if recovered: lines.append(f'Nomes recuperados de outras linhas da mesma barra em {len(recovered)} barras; a linha de referência fica registada.')
    return lines


def members(records):
    result = {}
    for r in records:
        m = result.setdefault(r['member'], {'support': r['support'], 'story': r['story'], 'nodes': set(), 'section': r['section'], 'section_issue': r['section_issue']})
        if any(m[k] != r[k] for k in ('support', 'story', 'section')): raise ValueError(f"Barra {r['member']}: identificação ou secção variável. Separe as revisões ou verifique o modelo.")
        m['nodes'].add(r['node'])
    return result


def all_records(tables):
    records = []; seen = set()
    for data in tables:
        for r in parse_bars(restore(data)):
            k = key(r['member'], r['node'], r['combination'])
            if k in seen: raise ValueError(f"Resultado Modelo repetido entre ficheiros: {r['member']}/{r['node']}/{r['combination']}.")
            seen.add(k); records.append(r)
    _resolve_names(records)
    members(records)
    return records


def parse_nodes(table):
    if profile(table.headers) != 'analysis_nodes': raise ValueError('Tabela de nós: exporte Node, X, Y e Z, com unidades nas coordenadas.')
    mapping = column_map(table); decimal = detect_decimal(table, mapping); result = {}
    for axis in ('x', 'y', 'z'):
        if mapping[axis][1] not in UNITS['hy']: raise ValueError(f'Coordenada {axis.upper()}: indique m, cm ou mm no cabeçalho.')
    for row, raw in table.rows:
        if any((row, mapping[k][0]) in table.formulas for k in ('node', 'x', 'y', 'z')): raise ValueError(f'Linha {row}: fórmula/erro XLSX nos nós.')
        node = identity(raw[mapping['node'][0]], 'Nó')
        if not node.isdigit() or int(node) <= 0: raise ValueError(f'Linha {row}: número de nó inválido.')
        if node in result: raise ValueError(f'Nó {node} repetido.')
        result[node] = {axis: _number(raw[mapping[axis][0]], axis.upper(), decimal)*UNITS['hy'][mapping[axis][1]] for axis in ('x', 'y', 'z')}
    return result


def prepare_analysis_table(table):
    mode = profile(table.headers)
    if mode == 'analysis': parse_bars(table)
    elif mode == 'analysis_nodes': parse_nodes(table)
    else: raise ValueError('Tabela de barras não reconhecida.')
    return {'mode': mode, 'geometries': {}, 'loads': {}, 'file': table.file, 'sha256': table.sha256,
            'row_count': len(table.rows), mode + '_table': archive(table)}


def member_role(member, node, index, nodes):
    """No floor-name or node-number heuristic. Missing coordinates remain unknown."""
    ends = index[member]['nodes']
    if len(ends) != 2 or node not in ends: return '', 'São necessários os dois nós de extremidade da barra.'
    other = next(x for x in ends if x != node)
    if node not in nodes or other not in nodes: return '', 'Cotas por confirmar; importe a tabela de nós ou confira os extremos no modelo.'
    a, b = nodes[node], nodes[other]
    if abs(a['x']-b['x']) > 1e-5 or abs(a['y']-b['y']) > 1e-5 or abs(a['z']-b['z']) < 1e-5:
        return 'unsupported', 'Barra não vertical: o equilíbrio deste perfil não se aplica.'
    return ('below' if a['z'] > b['z'] else 'above'), ''


def geometry_for(section, local_y):
    if section is None: return None
    hy, hz = section['hy'], section['hz']
    if section['shape'] == 'circular': return {'pilar_forma': 'circular', 'pilar_c1': hy, 'pilar_c2': None}
    return {'pilar_forma': 'retangular', 'pilar_c1': hy if local_y.endswith('X') else hz,
            'pilar_c2': hz if local_y.endswith('X') else hy}


def _validate_setup(setup, index, nodes):
    for label in ('floor', 'support', 'node', 'below'): identity(setup.get(label), label)
    for label in ('analysis_3d', 'same_combination', 'isolated_joint', 'axes_confirmed', 'at_joint'):
        if setup.get(label) is not True: raise ValueError('Confirme o âmbito, as combinações, a orientação dos eixos e os extremos na cota da laje.')
    if not str(setup.get('reference', '')).strip(): raise ValueError('Indique a referência do modelo e da conferência dos eixos.')
    if type(setup.get('upper_absent')) is not bool: raise ValueError('Declare explicitamente se o superior está fisicamente ausente.')
    if setup.get('above') == setup['below']: raise ValueError('Os tramos inferior e superior não podem ser a mesma barra.')
    if setup.get('above') and setup.get('upper_absent'): raise ValueError('Existe tramo superior e está declarado ausente.')
    extras = [m for m, data in index.items() if setup['node'] in data['nodes'] and m not in (setup['below'], setup.get('above'))]
    if extras: raise ValueError('Existem outras barras com resultados no nó (' + ', '.join(extras) + '). Associe os tramos corretos; se há outras transferências, use resultantes integrais da ligação.')
    for role in ('below', 'above'):
        m = setup.get(role)
        if not m: continue
        if m not in index or setup['node'] not in index[m]['nodes']: raise ValueError('As barras têm de partilhar o mesmo número de nó. Não se associam nós próximos automaticamente.')
        found, issue = member_role(m, setup['node'], index, nodes)
        if len(index[m]['nodes']) != 2: raise ValueError(issue)
        if found and found != role: raise ValueError(f'Barra {m}: posição incompatível com as coordenadas dos nós. {issue}')
        orient = setup.get(role + '_axes', {})
        if orient.get('x') not in ('+Z', '-Z') or orient.get('y') not in AXES:
            raise ValueError(f'Barra {m}: indique x local (+Z/-Z) e y local (+X/-X/+Y/-Y) relativamente aos eixos do caso.')


def build_joint(tables, setup, node_table=None):
    records = all_records(tables); index = members(records)
    nodes = parse_nodes(restore(node_table)) if node_table else {}
    _validate_setup(setup, index, nodes)
    chosen = [r for r in records if r['node'] == setup['node'] and r['member'] in (setup['below'], setup.get('above'))]
    combos = sorted({r['combination'] for r in chosen})
    if setup.get('combination'):
        if setup['combination'] not in combos: raise ValueError('Combinação não encontrada neste nó.')
        combos = [setup['combination']]
    loads = {}; geometries = {}
    # Keep node evidence bounded per connection, rather than duplicating a floor
    # table in every case. Coordinates of both member endpoints are sufficient.
    compact_nodes = None
    if node_table:
        ends = set().union(*(index[m]['nodes'] for m in (setup['below'], setup.get('above')) if m))
        nt = restore(node_table); col = column_map(nt)['node'][0]
        nt.rows = [(i, raw) for i, raw in nt.rows if identity(raw[col], 'Nó') in ends]
        compact_nodes = archive(nt)
    for combo in combos:
        compact = []
        # Current combination plus one row for each other endpoint, where that
        # combination is missing. This preserves topology without inventing loads.
        retained = {key(r['member'], r['node'], r['combination']) for r in records if r['member'] in (setup['below'], setup.get('above')) and r['combination'] == combo}
        endpoints = {key(r['member'], r['node']) for r in records if key(r['member'], r['node'], r['combination']) in retained}
        for r in records:
            e = key(r['member'], r['node'])
            if r['member'] in (setup['below'], setup.get('above')) and e not in endpoints:
                retained.add(key(r['member'], r['node'], r['combination'])); endpoints.add(e)
        # If a blank Name was resolved from another combination, archive that
        # exact source row too, so reopening reconstructs the same identity.
        for r in records:
            if key(r['member'], r['node'], r['combination']) not in retained: continue
            evidence = r.get('name_resolution', {}).get('evidence')
            if evidence:
                raw = evidence['raw']
                compound = next(value for h, value in raw.items() if header(h)[0] in COMPOUND)
                retained.add(key(*split_key(compound)))
        for data in tables:
            t = restore(data); col = column_map(t)['compound'][0]
            t.rows = [(i, raw) for i, raw in t.rows if key(*split_key(raw[col])) in retained]
            if t.rows: compact.append(archive(t))
        items = [deepcopy(r) for r in chosen if r['combination'] == combo]; issues = []
        b = next((r for r in items if r['member'] == setup['below']), None)
        a = next((r for r in items if r['member'] == setup.get('above')), None)
        if not b: issues.append('Faltam os esforços do tramo inferior para esta combinação.')
        if not a and not setup.get('upper_absent'): issues.append('Faltam os esforços do tramo superior; ausência de dados não declara ausência física.')
        if re.search(r'(?i)\b(max|min|envelope|envolvente|extrema)\b', combo): issues.append('Caso de envolvente: exporte combinações concomitantes individuais.')
        modal_issue = modal_combination_issue(combo)
        if modal_issue: issues.append(modal_issue)
        totals = {'V_Ed': 0., 'M_Edx': 0., 'M_Edy': 0.}
        for r in items:
            role = 'below' if r is b else 'above'; orient = setup[role + '_axes']
            xsign = 1 if orient['x'] == '+Z' else -1
            endpoint = 'end' if (role == 'below') == (xsign == 1) else 'begin'
            if modal_issue:
                r.update(role=role, endpoint=endpoint, joint_mx=None, joint_my=None, axes=deepcopy(orient))
                continue
            sign = 1 if endpoint == 'end' else -1
            yx, yy = AXES[orient['y']]; zx, zy = -xsign*yy, xsign*yx
            mx = sign*(r['values']['my']*yx + r['values']['mz']*zx)
            my = sign*(r['values']['my']*yy + r['values']['mz']*zy)
            totals['V_Ed'] += r['values']['fx'] * (1 if role == 'below' else -1)
            totals['M_Edx'] += mx; totals['M_Edy'] -= my
            r.update(role=role, endpoint=endpoint, joint_mx=mx, joint_my=my, axes=deepcopy(orient))
        if not modal_issue and totals['V_Ed'] <= 0: issues.append('VEd não positivo: reveja o equilíbrio e a face tracionada. Não se aplica valor absoluto.')
        if not modal_issue and not all(math.isfinite(v) for v in totals.values()): issues.append('Resultante não finita.')
        s = deepcopy(setup); s['combination'] = combo
        candidate = {'floor': setup['floor'].strip(), 'support': setup['support'].strip(), 'combination': combo,
            'kind': 'columns', 'adopted': None if issues else totals, 'issues': issues, 'records': items,
            'method': 'Modelo 3D: VEd=FXinf-FXsup; momentos locais convertidos em ações sobre o mesmo nó (+ no fim, - no início); MEdx=ΣMx,nó; MEdy=-ΣMy,nó.',
            'reference': setup['reference'].strip(), 'adapter': ADAPTER,
            'analysis': {'tables': compact, 'setup': s, 'nodes': compact_nodes}}
        if modal_issue: candidate['method'] = 'Modelo CQC/SRSS: esforços de origem conservados; resultante da ligação por obter com tratamento modal adequado.'
        loads[key(candidate['floor'], candidate['support'], combo)] = candidate
        section = index[setup['below']]['section']
        geometry = geometry_for(section, setup['below_axes']['y'])
        if geometry and b:
            geometries[key(candidate['floor'], candidate['support'])] = {'floor': candidate['floor'], 'support': candidate['support'], 'inputs': geometry, 'source': b['source']}
    if not loads: raise ValueError('Não existem casos no nó selecionado.')
    return {'mode': 'columns', 'loads': loads, 'geometries': geometries, 'file': 'Ligação Modelo · nó '+setup['node'], 'sha256': '', 'row_count': len(chosen)}


def validate_candidate(candidate):
    if candidate.get('adapter') == 'Modelo.joint-frame/1':
        from .model_workflow import validate_frame
        return validate_frame(candidate)
    if candidate.get('adapter') != ADAPTER: raise ValueError('Versão do adaptador Modelo desconhecida.')
    data = candidate['analysis']
    if modal_combination_issue(candidate['combination']) and candidate.get('adopted') is not None:
        raise ValueError('Esta ligação guardada adotou esforços CQC/SRSS por equilíbrio direto. Reimporte as tabelas e obtenha uma resultante da ligação com tratamento modal adequado.')
    batch = build_joint(data['tables'], data['setup'], data['nodes'])
    rebuilt = batch['loads'][key(candidate['floor'], candidate['support'], candidate['combination'])]
    if candidate != rebuilt: raise ValueError('Os esforços Modelo não correspondem à tabela e à associação guardadas.')


def analysis_source_lines(candidate):
    if candidate.get('adapter') == 'Modelo.joint-frame/1':
        from .model_workflow import frame_source_lines
        return frame_source_lines(candidate)
    setup = candidate['analysis']['setup']
    modal = bool(modal_combination_issue(candidate['combination']))
    lines = ['Modelo · esforços de barras', candidate['method'], f"Nó da ligação: {setup['node']}; inferior: {setup['below']}; superior: {setup.get('above') or ('ausente, declarado' if setup.get('upper_absent') else 'por associar')}.",
        'Âmbito declarado: barras 3D verticais alinhadas; esforços originais ELU concomitantes; extremos no centro do pilar na cota da laje, sem offsets; sem outras transferências nodais.',
        'X/Y são os eixos da planta do caso. A ordem das linhas e os números dos nós não definem início/fim.',
        'Convenção de sinais Modelo: ' + SIGN_REFERENCE]
    if modal: lines[3] = 'Concomitância não estabelecida pelos resultados CQC/SRSS; não foi calculada uma resultante por equilíbrio direto.'
    for r in candidate['records']:
        s = r['source']; v = r['values']
        label = 'Identificação' if r.get('name_resolution') else 'Name'
        axes = f"Eixos locais declarados: x={r['axes']['x']}, y={r['axes']['y']}; nó {'final' if r['endpoint']=='end' else 'inicial'}."
        if modal: axes += ' Ações concomitantes sobre o nó por determinar.'
        else: axes += f" Ações sobre o nó: Mx={r['joint_mx']:.6g}, My={r['joint_my']:.6g} kN.m."
        lines.extend([f"Barra {r['member']}, nó {r['node']}, {r['combination']}; Story={r['story']}; {label}={r['support']}.",
            f"FX={v['fx']:.6g} kN (compressão positiva); MY={v['my']:.6g}, MZ={v['mz']:.6g} kN.m (locais).",
            axes,
            f"Origem: {s['file']}, folha {s['sheet']}, linha {s['row']}; SHA-256 {s['sha256']}.",
            'Valores originais e unidades: ' + '; '.join(f'{k}={v}' for k, v in s['raw'].items())])
        if name_note(r): lines.append(name_note(r))
    if candidate['analysis']['nodes']:
        t = candidate['analysis']['nodes']; lines.append(f"Coordenadas conferidas com {t['file']}, folha {t['sheet']}; SHA-256 {t['sha256']}.")
    else: lines.append('Sem tabela de nós: posição dos tramos e cota da ligação conferidas pelo utilizador.')
    lines.append('Referência da conferência: ' + candidate['reference'])
    lines.extend('Por resolver: '+x for x in candidate['issues'])
    if candidate['adopted']: lines.append('Resultante: ' + '; '.join(f'{k}={v:.6g}' for k, v in candidate['adopted'].items()) + ' (kN; kN.m).')
    return lines
