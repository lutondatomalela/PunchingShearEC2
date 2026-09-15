"""Read external tables and derive connection actions, independently of EC2.

All normalized forces are kN and kN.m; geometry is in metres. Column moments
are actions ON the joint in a right-handed X/Y/Z frame, Z upward. The engine
uses ex=MEdy/V and ey=MEdx/V: MEdx=sum(Mx_joint), MEdy=-sum(My_joint).
No member-local transformation, envelope pairing or FE integration is inferred.
"""
from copy import deepcopy
import csv
from dataclasses import dataclass, asdict
import hashlib
import io
import json
import math
from pathlib import Path
import re
import unicodedata


MODES = {'geometry': 'Geometria dos pilares', 'columns': 'Esforços nos extremos dos pilares',
         'resultants': 'Resultantes integrais da ligação',
         'analysis': 'Modelo · esforços de barras (reconhecimento automático)',
         'analysis_nodes': 'Modelo · coordenadas dos nós (reconhecimento automático)'}
FIELDS = {
    'floor': 'Piso', 'support': 'Pilar / ligação', 'combination': 'Combinação ELU',
    'shape': 'Forma: retangular / circular', 'c1': 'c1 / diâmetro', 'c2': 'c2',
    'member': 'Barra', 'role': 'Tramo: inferior / superior', 'end': 'Extremo: topo / base',
    'upper_absent': 'Sem tramo superior: sim / não', 'n': 'N axial',
    'mx': 'Mx no nó', 'my': 'My no nó', 'v': 'V da ligação',
}
MODE_FIELDS = {
    'geometry': ('floor', 'support', 'shape', 'c1', 'c2'),
    'columns': ('floor', 'support', 'combination', 'member', 'role', 'end', 'upper_absent', 'n', 'mx', 'my', 'shape', 'c1', 'c2'),
    'resultants': ('floor', 'support', 'combination', 'v', 'mx', 'my', 'shape', 'c1', 'c2'),
}
REQUIRED = {'geometry': {'floor', 'support', 'shape', 'c1'},
            'columns': {'floor', 'support', 'combination', 'member', 'role', 'end', 'n', 'mx', 'my'},
            'resultants': {'floor', 'support', 'combination', 'v', 'mx', 'my'}}
ALIASES = {
    'floor': ('floor', 'piso', 'nivel'), 'support': ('support', 'pilar', 'ligacao'),
    'combination': ('combination', 'combinacao', 'caso'), 'shape': ('shape', 'forma'),
    'c1': ('c1', 'diametro', 'c1_mm', 'c1_m'), 'c2': ('c2', 'c2_mm', 'c2_m'),
    'member': ('member', 'barra'), 'role': ('role', 'tramo'), 'end': ('end', 'extremo'),
    'upper_absent': ('upper_absent', 'sem_superior'), 'n': ('n', 'n_kn', 'n_n'),
    'mx': ('mx', 'mx_no', 'mx_knm', 'medx'), 'my': ('my', 'my_no', 'my_knm', 'medy'),
    'v': ('v', 'v_kn', 'ved'),
}
LENGTH = {'m': 1., 'cm': .01, 'mm': .001}
FORCE = {'kN': 1., 'N': .001}
MOMENT = {'kN.m': 1., 'N.m': .001, 'N.mm': .000001, 'kN.cm': .01}
MAX_ROWS = 50000


def token(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value).strip().casefold()) if not unicodedata.combining(c))


def identity(value, label):
    if value is None or isinstance(value, bool): raise ValueError(f'{label}: identificação em falta.')
    if isinstance(value, float):
        if not math.isfinite(value): raise ValueError(f'{label}: identificação inválida.')
        value = int(value) if value.is_integer() else value
    result = str(value).strip()
    if not result: raise ValueError(f'{label}: identificação em falta.')
    if len(result) > 200: raise ValueError(f'{label}: identificação demasiado longa.')
    return result


def numeric(value, label, decimal='.'):
    if isinstance(value, bool) or value is None: raise ValueError(f'{label}: valor em falta ou inválido.')
    if isinstance(value, (float, int)): number = float(value)
    else:
        text = str(value).strip().replace('\u2212', '-')
        if decimal not in ('.', ','): raise ValueError('Separador decimal inválido.')
        pattern = r'[+-]?(?:\d+(?:' + re.escape(decimal) + r'\d*)?|' + re.escape(decimal) + r'\d+)(?:[eE][+-]?\d+)?'
        if not re.fullmatch(pattern, text):
            raise ValueError(f'{label}: use decimal «{decimal}», sem separador de milhares; recebido {text!r}.')
        number = float(text.replace(',', '.'))
    if not math.isfinite(number): raise ValueError(f'{label}: o valor tem de ser finito.')
    return number


def boolean(value):
    if value is None or str(value).strip() == '': return False
    if token(value) in ('sim', 'yes', 'true', '1'): return True
    if token(value) in ('nao', 'no', 'false', '0'): return False
    raise ValueError('Sem tramo superior: use sim ou não.')


@dataclass
class Table:
    file: str
    sha256: str
    sheet: str
    headers: list
    rows: list  # [(physical row number, values)]
    formulas: set


def worksheets(path):
    if Path(path).suffix.lower() != '.xlsx': return ['CSV']
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False)
    try: return wb.sheetnames
    finally: wb.close()


def read_table(path, sheet=None, header_row=1, delimiter='auto', encoding='auto'):
    path = Path(path)
    if path.stat().st_size > 32 * 1024 * 1024: raise ValueError('A tabela excede 32 MB. Exporte apenas os dados necessários.')
    if type(header_row) is not int or not 1 <= header_row <= 500: raise ValueError('Linha de cabeçalhos: indique 1 a 500.')
    payload = path.read_bytes(); digest = hashlib.sha256(payload).hexdigest(); formulas = set()
    records = []
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            name = sheet or wb.sheetnames[0]
            if name not in wb.sheetnames: raise ValueError('Folha não encontrada.')
            ws = wb[name]
            # Some valid exporters omit dimensions; others write A1 for a larger
            # table. Read actual XML rows and enforce bounds during streaming.
            ws.reset_dimensions()
            for cells in ws.iter_rows(min_row=header_row):
                if len(cells)>256 or len(records)>MAX_ROWS:
                    raise ValueError('Exporte uma tabela até 256 colunas e 50 000 linhas de dados.')
                values = []
                for col, cell in enumerate(cells):
                    value = cell.value
                    if cell.data_type in ('f', 'e'): formulas.add((len(records) + header_row, col))
                    # Preserve numeric identifiers explicitly formatted with leading zeros.
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and re.fullmatch(r'0{2,}', cell.number_format or ''):
                        if math.isfinite(value) and float(value).is_integer(): value = str(int(value)).zfill(len(cell.number_format))
                    if value is not None and not isinstance(value, (str, int, float, bool)): value = str(value)
                    values.append(value)
                records.append(values)
        finally: wb.close()
    elif path.suffix.lower() in ('.csv', '.tsv', '.txt'):
        if encoding == 'auto': encoding = 'utf-16' if payload.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        try: contents = payload.decode(encoding)
        except UnicodeError as exc: raise ValueError('Codificação não reconhecida. Selecione a codificação do ficheiro.') from exc
        if delimiter == 'auto':
            try: delimiter = csv.Sniffer().sniff('\n'.join(contents.splitlines()[header_row-1:header_row+20]), delimiters=';\t,').delimiter
            except csv.Error as exc: raise ValueError('Escolha o separador de colunas: ponto e vírgula, tabulação ou vírgula.') from exc
        if delimiter not in (';', ',', '\t'): raise ValueError('Separador de colunas inválido.')
        records = list(csv.reader(io.StringIO(contents), delimiter=delimiter))
        records = records[header_row-1:]; name = 'CSV'
    else: raise ValueError('Formato suportado: XLSX, CSV ou TSV. Converta os ficheiros XLS em XLSX.')
    if not records: raise ValueError('Tabela vazia.')
    if len(records) > MAX_ROWS + 1: raise ValueError('A tabela excede 50 000 linhas de dados.')
    while records[0] and records[0][-1] in (None, '') and all(len(r) < len(records[0]) or r[len(records[0])-1] in (None, '') for r in records): records[0].pop()
    headers = [str(x).strip() if x is not None else '' for x in records[0]]
    if not headers or len(headers) > 256 or any(not h for h in headers): raise ValueError('Os cabeçalhos têm de estar preenchidos (até 256 colunas).')
    if len({token(h) for h in headers}) != len(headers): raise ValueError('Existem cabeçalhos repetidos. Dê um nome distinto a cada coluna.')
    rows = []
    for i, values in enumerate(records[1:], header_row+1):
        if not any(v is not None and str(v).strip() for v in values): continue
        if any(v not in (None, '') for v in values[len(headers):]): raise ValueError(f'Linha {i}: existem dados sem cabeçalho.')
        rows.append((i, values[:len(headers)] + [None] * max(0, len(headers)-len(values))))
    if not rows: raise ValueError('Não existem linhas de dados.')
    return Table(path.name, digest, name, headers, rows, formulas)


@dataclass
class ImportConfig:
    mode: str = 'geometry'
    length_unit: str = 'm'
    force_unit: str = 'kN'
    moment_unit: str = 'kN.m'
    decimal: str = '.'
    compression: str = 'positive'
    resultant_convention: str = 'program'
    same_combination: bool = False
    axes_at_column: bool = False
    isolated_joint: bool = False
    full_connection: bool = False
    reference: str = ''

    def validate(self):
        if self.mode not in MODES: raise ValueError('Tipo de tabela desconhecido.')
        if self.length_unit not in LENGTH or self.force_unit not in FORCE or self.moment_unit not in MOMENT: raise ValueError('Unidades não suportadas.')
        if self.decimal not in ('.', ',') or self.compression not in ('positive', 'negative') or self.resultant_convention not in ('program', 'vector'): raise ValueError('Convenção de importação inválida.')
        if self.mode in ('columns', 'resultants'):
            if self.same_combination is not True: raise ValueError('Confirme esforços originais de análise da mesma combinação ELU, sem envolventes independentes.')
            if self.axes_at_column is not True: raise ValueError('Confirme os eixos e a redução dos esforços ao centro do pilar, na cota da laje.')
        if self.mode == 'columns' and self.isolated_joint is not True:
            raise ValueError('O equilíbrio automático requer tramos verticais alinhados e apenas a laje a transferir carga ao nó.')
        if self.mode == 'resultants' and self.full_connection is not True:
            raise ValueError('Confirme resultantes da ligação completa, sem redução por cargas interiores ao contorno e sem majoração por beta.')


def suggest_mapping(table, mode):
    if mode in ('analysis', 'analysis_nodes'): return {}
    normalized = {token(h): h for h in table.headers}
    return {k: next((normalized[a] for a in ALIASES[k] if a in normalized), '') for k in MODE_FIELDS[mode]}


def key(*parts): return json.dumps(parts, ensure_ascii=False, separators=(',', ':'))


def _geometry(values, scale, decimal):
    if not any(values.get(k) not in (None, '') for k in ('shape', 'c1', 'c2')): return None
    shape = token(values.get('shape', ''))
    shape = {'rectangular': 'retangular', 'rectangle': 'retangular', 'circle': 'circular'}.get(shape, shape)
    if shape not in ('retangular', 'circular'): raise ValueError('Forma: indique retangular ou circular.')
    c1 = numeric(values.get('c1'), 'c1 / D', decimal) * scale
    c2 = numeric(values.get('c2'), 'c2', decimal) * scale if shape == 'retangular' else None
    if not .01 <= c1 <= 20 or (c2 is not None and not .01 <= c2 <= 20): raise ValueError('Dimensões fora de 0,01 a 20 m. Confira as unidades.')
    if shape == 'circular' and values.get('c2') not in (None, ''):
        if not math.isclose(numeric(values['c2'], 'c2', decimal)*scale, c1): raise ValueError('Pilar circular: c2 deve estar vazio ou ser igual ao diâmetro.')
    return {'pilar_forma': shape, 'pilar_c1': c1, 'pilar_c2': c2}


def prepare_import(table, mapping, config):
    """Atomic conversion. Invalid rows are never dropped or defaulted to zero."""
    config.validate()
    if config.mode in ('analysis', 'analysis_nodes'):
        from .analysis_import import prepare_analysis_table
        batch = prepare_analysis_table(table)
        if batch['mode'] != config.mode: raise ValueError('O tipo selecionado não corresponde à tabela de barras.')
        return batch
    mapping = {k: h for k, h in mapping.items() if h and k in MODE_FIELDS[config.mode]}
    missing = REQUIRED[config.mode] - mapping.keys()
    if missing: raise ValueError('Mapeie os campos obrigatórios: ' + ', '.join(FIELDS[k] for k in sorted(missing)))
    if any(h not in table.headers for h in mapping.values()): raise ValueError('O mapeamento refere cabeçalhos inexistentes.')
    if len(set(mapping.values())) != len(mapping): raise ValueError('Cada coluna de origem só pode ser associada a uma grandeza.')
    indexes = {k: table.headers.index(h) for k, h in mapping.items()}
    groups = {}; geometries = {}; errors = []
    for row_number, raw in table.rows:
        try:
            if any((row_number, i) in table.formulas for i in indexes.values()):
                raise ValueError('Célula com fórmula/erro XLSX. Exporte os valores calculados antes de importar.')
            values = {k: raw[i] for k, i in indexes.items()}
            floor, support = (identity(values.get(k), FIELDS[k]) for k in ('floor', 'support'))
            source = {'file': table.file, 'sha256': table.sha256, 'sheet': table.sheet, 'row': row_number,
                      'raw': dict(zip(table.headers, raw)), 'mapping': mapping, 'config': asdict(config)}
            geometry = _geometry(values, LENGTH[config.length_unit], config.decimal)
            if config.mode == 'geometry':
                if geometry is None: raise ValueError('Dimensões e forma do pilar em falta.')
                gkey = key(floor, support)
                if gkey in geometries: raise ValueError('Piso/pilar repetido na tabela de geometria.')
                geometries[gkey] = {'floor': floor, 'support': support, 'inputs': geometry, 'source': source}
                continue
            combo = identity(values.get('combination'), FIELDS['combination'])
            item = {'floor': floor, 'support': support, 'combination': combo, 'geometry': geometry, 'source': source}
            for k in ('mx', 'my', 'n' if config.mode == 'columns' else 'v'):
                item[k] = numeric(values.get(k), FIELDS[k], config.decimal) * (FORCE[config.force_unit] if k in ('n', 'v') else MOMENT[config.moment_unit])
            if config.mode == 'columns':
                role = token(values.get('role')); end = token(values.get('end'))
                item['role'] = {'inferior': 'below', 'superior': 'above', 'below': 'below', 'above': 'above'}.get(role)
                item['end'] = {'topo': 'top', 'base': 'bottom', 'top': 'top', 'bottom': 'bottom'}.get(end)
                if not item['role'] or not item['end']: raise ValueError('Tramo: inferior/superior. Extremo: topo/base.')
                if (item['role'], item['end']) not in (('below', 'top'), ('above', 'bottom')):
                    raise ValueError('Use o topo do tramo inferior e a base do tramo superior, junto à mesma laje.')
                item['member'] = identity(values.get('member'), 'Barra')
                item['upper_absent'] = boolean(values.get('upper_absent'))
                item['n'] *= 1 if config.compression == 'positive' else -1
            groups.setdefault(key(floor, support, combo), []).append(item)
        except (ValueError, TypeError) as exc: errors.append(f'Linha {row_number}: {exc}')
    if errors: raise ValueError('\n'.join(errors[:20]) + (f'\nMais {len(errors)-20} erros.' if len(errors)>20 else ''))
    loads = {}
    for ckey, items in groups.items():
        first = items[0]; issues = []; adopted = None
        if config.mode == 'resultants':
            if len(items) != 1: issues.append('Mais de uma resultante para o mesmo piso, pilar e combinação. Não somar cortes ou envolventes automaticamente.')
            else: adopted = {'V_Ed': first['v'], 'M_Edx': first['mx'], 'M_Edy': first['my'] * (1 if config.resultant_convention == 'program' else -1)}
            method = 'Resultante integral no centro do pilar; VEd=V; MEdx=Mx; MEdy=' + ('My.' if config.resultant_convention == 'program' else '-My (convenção vetorial).')
            geometry = first['geometry']
        else:
            below = [x for x in items if x['role']=='below']; above = [x for x in items if x['role']=='above']
            if len(below) != 1 or len(above) > 1: issues.append('É necessário exatamente um tramo inferior e, quando existente, um superior por combinação.')
            if any(x['upper_absent'] for x in above): issues.append('A declaração de ausência do superior pertence à linha do tramo inferior.')
            if len(below) == 1:
                if not above and not below[0]['upper_absent']: issues.append('Tramo superior em falta. Falta de dados não equivale a ausência física.')
                if above and below[0]['upper_absent']: issues.append('O superior está declarado ausente mas existem esforços desse tramo.')
            if len({x['member'] for x in items}) != len(items): issues.append('A mesma barra surge mais de uma vez na ligação/combinação.')
            if not issues:
                b = below[0]; a = above[0] if above else {'n': 0, 'mx': 0, 'my': 0}
                adopted = {'V_Ed': b['n']-a['n'], 'M_Edx': b['mx']+a['mx'], 'M_Edy': -(b['my']+a['my'])}
            geometry = below[0]['geometry'] if len(below)==1 else None
            method = 'Equilíbrio no nó: VEd=Ninf-Nsup; MEdx=Mx,inf+Mx,sup; MEdy=-(My,inf+My,sup). N normalizado positivo em compressão; momentos vetoriais de ações dos pilares sobre o nó.'
        if adopted is not None and not all(math.isfinite(v) for v in adopted.values()):
            issues.append('Resultante não finita. Reveja os valores e as unidades.'); adopted = None
        if adopted is not None and adopted['V_Ed'] <= 0:
            issues.append('VEd não positivo. Não se aplica valor absoluto: reveja sinais, equilíbrio e face tracionada da laje.'); adopted = None
        if issues: adopted = None
        loads[ckey] = {'floor': first['floor'], 'support': first['support'], 'combination': first['combination'],
                       'kind': config.mode, 'adopted': adopted, 'method': method, 'issues': issues,
                       'records': items, 'reference': config.reference.strip()}
        if geometry:
            gkey = key(first['floor'], first['support'])
            source = next(x['source'] for x in items if x['geometry'] == geometry)
            geom = {'floor': first['floor'], 'support': first['support'], 'inputs': geometry, 'source': source}
            if gkey in geometries and geometries[gkey]['inputs'] != geometry:
                raise ValueError('Geometrias diferentes para o mesmo pilar/piso entre combinações.')
            geometries[gkey] = geom
    return {'mode': config.mode, 'geometries': geometries, 'loads': loads,
            'file': table.file, 'sha256': table.sha256, 'row_count': len(table.rows)}


def source_lines(candidate):
    if candidate.get('adapter'):
        from .analysis_import import analysis_source_lines
        return analysis_source_lines(candidate)
    lines = [MODES[candidate['kind']], candidate['method']]
    lines.append('Âmbito declarado: esforços originais de análise e concomitantes, reduzidos ao centro do pilar na cota da laje; eixos compatíveis com a planta.')
    if candidate['kind'] == 'columns':
        lines.append('Tramos verticais alinhados; a laje é a única origem da transferência no nó, sem outras cargas nodais. Ninf no topo do inferior e Nsup na base do superior.')
    else:
        lines.append('Resultante da ligação completa, sem dedução de cargas interiores a um contorno nem majoração por beta.')
    for item in candidate['records']:
        s = item['source']; units = s['config']
        lines.append(f"Origem: {s['file']}, folha {s['sheet']}, linha {s['row']}. SHA-256: {s['sha256']}")
        lines.append('Colunas: ' + '; '.join(f"{FIELDS[k]}={h}" for k, h in s['mapping'].items()))
        lines.append('Valores originais: ' + '; '.join(f'{h}={s["raw"].get(h)}' for h in s['mapping'].values()))
        lines.append(f"Unidades: {units['length_unit']}; {units['force_unit']}; {units['moment_unit']}. Decimal: {units['decimal']}.")
        if candidate['kind'] == 'columns':
            lines.append('Sinal de N no ficheiro: compressão ' + ('positiva.' if units['compression']=='positive' else 'negativa.'))
            lines.append(f"Tramo {'inferior, topo' if item['role']=='below' else 'superior, base'}; barra {item['member']}: N={item['n']:.6g} kN (compressão positiva), Mx,nó={item['mx']:.6g} kN.m, My,nó={item['my']:.6g} kN.m. Superior ausente: {'sim' if item['upper_absent'] else 'não'}.")
    if candidate['reference']: lines.append('Referência: ' + candidate['reference'])
    if candidate['adopted']: lines.append('Resultante obtida: ' + '; '.join(f'{k}={v:.6g}' for k, v in candidate['adopted'].items()) + ' (kN; kN.m).')
    lines.extend('Por resolver: ' + issue for issue in candidate['issues'])
    return lines


def comparison_lines(candidates):
    if not all(candidates.get(k, {}).get('adopted') for k in ('columns', 'resultants')): return []
    a = candidates['columns']['adopted']; b = candidates['resultants']['adopted']
    return ['Comparação de fontes (resultante da ligação menos equilíbrio dos pilares; não é critério automático de aceitação): ' +
            '; '.join(f'{k}: {b[k]-a[k]:+.6g}' for k in a) + ' (kN; kN.m).']
