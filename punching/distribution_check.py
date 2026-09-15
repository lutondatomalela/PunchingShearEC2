"""Ensaio do executável real: recursos, Tk, cálculo, persistência e relatórios."""
import json
from pathlib import Path
import traceback

def run(output):
    target = Path(output).resolve()
    target.mkdir(parents=True, exist_ok=True)
    app = None
    evidence = {"status": "failed", "checks": []}
    try:
        import Punching_EC2_GUI as gui
        from .version import VERSION
        from .core import PuncoamentoEC2
        from .connections import ConnectionBook
        from .connection_ui import write_json_atomic
        from .reports import export_json, export_txt, export_pdf, export_xlsx
        from .table_import import read_table
        from openpyxl import load_workbook
        root = Path(gui.__file__).resolve().parent
        evidence["version"] = VERSION
        catalog = json.loads((root / "examples/catalog.json").read_text(encoding="utf-8"))
        assert len(catalog) == len(gui.EXAMPLES) and len(catalog) > 10
        evidence["checks"].append("catalogue")
        def error(*args, **kwargs):
            raise RuntimeError(str(args))
        gui.messagebox.showerror = error
        app = gui.PuncoamentoApp()
        app.report_callback_exception = lambda _type, value, tb: error(value)
        app.withdraw()
        app.update()
        entry = next(e for e in catalog if e["file"] == "19_interior_flexao_biaxial_643.json")
        app.var_exemplo.set(entry["title"])
        app.carregar_exemplo()
        app.update_idletasks()
        app.calcular()
        assert app.last_verif is not None and not app.dirty
        result = app.last_verif.snapshot()
        data = json.loads((root / "examples" / entry["file"]).read_text(encoding="utf-8"))
        reference = PuncoamentoEC2(**data)
        reference.verificar_puncoamento()
        for field in ("beta", "u1_eff", "v_Rd_c"):
            assert abs(result["values"][field] - reference.snapshot()["values"][field]) < 1e-10
        assert result["checks"]
        evidence["checks"].extend(["native_tk_startup", "gui_calculation"])
        # Exercise the GUI's actual save/reopen path with accented file names.
        case = target / "ligação.json"
        gui.filedialog.asksaveasfilename = lambda **kwargs: str(case)
        gui.filedialog.askopenfilename = lambda **kwargs: str(case)
        app.guardar_caso()
        app.limpar()
        app.abrir_caso()
        app.calcular()
        assert app.last_verif.snapshot()["inputs"] == result["inputs"]
        evidence["checks"].append("case_save_reopen")
        book = ConnectionBook.from_payload(json.loads(
            (root / "examples/importacao/Modelo_Ligacoes.json").read_text(encoding="utf-8")), gui.DEFAULTS)
        project = target / "conjunto.json"
        write_json_atomic(project, book.payload())
        restored = ConnectionBook.from_payload(json.loads(project.read_text(encoding="utf-8")), gui.DEFAULTS)
        assert len(restored.cases) == len(book.cases) > 0
        evidence["checks"].append("collection_save_reopen")
        for ext, exporter in (("json", export_json), ("txt", export_txt), ("pdf", export_pdf), ("xlsx", export_xlsx)):
            exporter(result, target / ("resultado." + ext))
        assert json.loads((target / "resultado.json").read_text(encoding="utf-8"))["values"] == result["values"]
        assert (target / "resultado.pdf").read_bytes().startswith(b"%PDF-")
        assert "Pun" in (target / "resultado.txt").read_text(encoding="utf-8")
        wb = load_workbook(target / "resultado.xlsx", read_only=True)
        assert "Entradas" in wb.sheetnames
        wb.close()
        evidence["checks"].append("exports_pdf_xlsx_txt_json")
        assert read_table(root / "examples/importacao/Exemplo_importacao.xlsx").rows
        assert read_table(root / "examples/importacao/Modelo_Nos.csv").rows
        evidence["checks"].append("imports_xlsx_utf16_csv")
        evidence["status"] = "passed"
    except Exception:
        evidence["error"] = traceback.format_exc()
    finally:
        if app is not None:
            app.destroy()
        (target / "self-test.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0 if evidence["status"] == "passed" else 1
