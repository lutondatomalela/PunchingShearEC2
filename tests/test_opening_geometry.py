"""Independent references for physical openings and their integration into EC2."""
import copy
import math
import json
from pathlib import Path
import pytest
from punching.openings import derive_openings,normalize_openings,placed_opening,dragged_opening,point_clearance
from punching.core import PuncoamentoEC2
from punching import geometry as geo
from punching.reports import export_xlsx,export_pdf,text_report
from test_reference_cases import run,inputs

CTX=dict(c1=.4,c2=.4,shape='retangular',d=.2,position='interior',edge_distance=0)

def rectangle(**kw):return dict(id='A1',shape='rectangular',x=1.,y=0.,width=.4,height=.4,**kw)

def derive(items,**kw):return derive_openings(items,**dict(CTX,**kw))

@pytest.mark.parametrize('side,centre',[('+X',(1,0)),('-X',(-1,0)),('+Y',(0,1)),('-Y',(0,-1))])
def test_square_tangents_at_each_face(side,centre):
    o=placed_opening('A1','rectangular',side,.6,.4,.4)
    assert (o['x'],o['y'])==centre
    _,records,sectors=derive([o]);r=records[0];alpha=math.degrees(math.atan(.2/.8))
    assert r['distance_m']==pytest.approx(.6)
    assert sectors[0][1]-sectors[0][0]==pytest.approx(2*alpha)
    angle=(sectors[0][1]+sectors[0][0])/2
    assert math.cos(math.radians(angle))==pytest.approx(centre[0],abs=1e-12)
    assert math.sin(math.radians(angle))==pytest.approx(centre[1],abs=1e-12)
    assert not r['equivalent_outline']
    for p,a in zip(r['tangent_points'],sectors[0]):assert math.atan2(p[1],p[0])==pytest.approx(math.atan2(math.sin(math.radians(a)),math.cos(math.radians(a))))

@pytest.mark.parametrize('gap,active',[(1.2-1e-7,True),(1.2,True),(1.2+1e-7,False)])
def test_exact_6d_distance_from_faces(gap,active):
    o=placed_opening('A1','rectangular','+X',gap,.4,.4)
    _,records,sectors=derive([o]);assert records[0]['active'] is active
    assert bool(sectors) is active

def test_diagonal_distance_uses_nearest_corners_not_centres():
    o=dict(rectangle(),x=1.,y=.7)
    _,records,_=derive([o]);assert records[0]['distance_m']==pytest.approx(math.hypot(.6,.3))

@pytest.mark.parametrize('shape',['retangular','circular'])
def test_circle_tangents_are_exact_and_not_a_sampled_polygon(shape):
    o=dict(id='C1',shape='circular',x=1.,y=0.,diameter=.4)
    _,records,sectors=derive([o],shape=shape)
    alpha=math.degrees(math.asin(.2))
    assert sectors[0]==pytest.approx([-alpha,alpha]);assert records[0]['distance_m']==pytest.approx(.6)
    for p in records[0]['tangent_points']:
        assert math.hypot(p[0]-1,p[1])==pytest.approx(.2)
        assert p[0]*(p[0]-1)+p[1]*p[1]==pytest.approx(0,abs=1e-12)

def test_elongated_radial_opening_uses_figure_614_equivalent_width():
    o=dict(rectangle(),x=1.,width=.8,height=.2)
    _,records,sectors=derive([o]);r=records[0]
    assert r['equivalent_width_m']==pytest.approx(.4)
    assert sectors[0]==pytest.approx([-math.degrees(math.atan(.2/.6)),math.degrees(math.atan(.2/.6))])
    assert r['sector_deg'][1]>r['tangent_sector_deg'][1]
    assert len(r['equivalent_outline'])==5

def test_oblique_equivalent_envelope_is_explicit_and_encloses_real_corners():
    o=dict(rectangle(),x=1.2,y=.3,width=.8,height=.2)
    _,records,_=derive([o]);r=records[0]
    assert 'hipótese geométrica adicional' in r['method']
    assert r['sector_deg'][0]<=r['tangent_sector_deg'][0]
    assert r['sector_deg'][1]>=r['tangent_sector_deg'][1]

@pytest.mark.parametrize('items,kw',[
    ([dict(rectangle(),x=.4)],{}),
    ([dict(id='C',shape='circular',x=.39,y=0,diameter=.4)],{}),
    ([dict(rectangle(),y=-.01)],dict(position='bordo')),
    ([dict(rectangle(),x=-.01,y=1)],dict(position='canto')),
    ([rectangle(),dict(rectangle(),id='A2',x=1.1)],{}),
    ([rectangle(),dict(rectangle(),id='A2',x=1.4)],{}),
    ([dict(rectangle(),x=1.,y=.4,width=4.,height=.2)],{}),
])
def test_contact_overlap_or_unsupported_envelope_is_rejected(items,kw):
    with pytest.raises(ValueError):derive(items,**kw)

@pytest.mark.parametrize('item',[
    dict(rectangle(),x=float('nan')),dict(rectangle(),width=0),dict(rectangle(),width=True),
    dict(rectangle(),rotation=30),dict(rectangle(),id=''),dict(rectangle(),shape='triangle')])
def test_invalid_physical_parameters_are_not_silently_ignored(item):
    with pytest.raises(ValueError):normalize_openings([item])

def test_manual_and_automatic_sectors_are_unioned_without_double_deduction():
    v=run(openings=[rectangle()],opening_sectors=[(-20,20)])
    manual=run(opening_sectors=[(-20,20)])
    assert v.status==manual.status=='BETA_PENDING'
    assert v.u1_eff==pytest.approx(manual.u1_eff)
    assert v.beta is None and not v.checks

def test_physical_square_matches_explicit_angles_and_keeps_original_inputs():
    o=rectangle();saved=copy.deepcopy(o)
    p=inputs(openings=[o]);v=PuncoamentoEC2(**p);v.verificar_puncoamento()
    alpha=math.degrees(math.atan(.2/.8));manual=run(opening_sectors=[(-alpha,alpha)])
    assert v.u1_eff==pytest.approx(manual.u1_eff)
    # Independently: u1 right face x=.60; rays y=+/-x*.25 remove .30 m.
    assert v.u1_eff==pytest.approx(1.6+4*math.pi*.2-.3)
    assert v.snapshot()['inputs']['opening_sectors'] is None
    assert o==saved
    r=v.snapshot();r['opening_records'][0]['sector_deg'][0]=999
    assert v.snapshot()['opening_records'][0]['sector_deg'][0]!=999

def test_far_opening_has_no_deduction_and_depth_change_recalculates_scope():
    o=dict(rectangle(),x=2.)
    small=run(openings=[o]);big=run(openings=[o],laje_d=.3)
    assert small.status=='PASS_WITHOUT' and not small.sectors
    assert big.status=='BETA_PENDING' and big.sectors

def test_simplified_and_anchorage_do_not_override_opening_scope():
    v=run(openings=[rectangle()],anchorage_confirmed=True,allow_biaxial_envelope=True)
    assert v.status=='BETA_PENDING'
    with pytest.raises(ValueError):PuncoamentoEC2(**inputs(openings=[rectangle()],beta_mode='simplificado',simplified_applicable=True))

def test_physical_opening_checks_actual_bar_clearance():
    v=run(openings=[rectangle()],V_Ed=450000,beta_mode='manual',beta_manual=1.2,beta_reference='Ensaio geométrico com fator prescrito.',anchorage_confirmed=True)
    assert v.status=='PASS_WITH'
    checks=[c for c in v.detail_checks if c['id']=='opening_clearance'];assert checks and all(c['passed'] for c in checks)
    # A corrupted bar must be caught independently of row counts/areas/flags.
    v.reinforcement_rows[0]['coordinates'].append((1.,0.))
    v._check_opening_clearance();assert v.status=='FAIL_DETAIL'

@pytest.mark.parametrize('shape',['rectangular','circular'])
def test_drag_operations_are_isolated_and_respect_dimensions(shape):
    o=placed_opening('A','rectangular' if shape=='rectangular' else 'circular','+X',.6,.4,.4)
    old=copy.deepcopy(o)
    moved=dragged_opening(o,'move',(1,0),(1.128,.174))
    assert moved['x']==pytest.approx(1.13) and moved['y']==pytest.approx(.17)
    bigger=dragged_opening(o,'resize',(1,0),(1.3,.4))
    if shape=='circular':assert bigger['diameter']==pytest.approx(1)
    else:assert (bigger['width'],bigger['height'])==pytest.approx((.6,.8))
    assert o==old

def test_reports_keep_measures_scope_and_exact_angles(tmp_path):
    v=run(openings=[rectangle()]);r=v.snapshot();text=text_report(r)
    assert 'Aberturas - definição geométrica' in text and '0.600' in text and 'BETA_PENDING' in text
    export_xlsx(r,tmp_path/'opening.xlsx');export_pdf(r,tmp_path/'opening.pdf')
    from openpyxl import load_workbook
    ws=load_workbook(tmp_path/'opening.xlsx')['Aberturas']
    assert ws['A2'].value=='A1' and ws['H2'].value==pytest.approx(.6)
    assert ws['K2'].value==pytest.approx(r['opening_records'][0]['sector_deg'][0])
