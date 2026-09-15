"""Regressões de coerência: controladores Tk sem janela e ficheiros exportados.

O interpretador Tcl e as variáveis/temporizadores são reais; os widgets visuais
e diálogos são substituídos. Estes testes não constituem ensaio visual da GUI.
"""
import json
import subprocess
import sys
from pathlib import Path
import tkinter as tk
import pytest
import Punching_EC2_GUI as gui
from punching.core import PuncoamentoEC2
from punching.reports import export_json, export_xlsx, export_pdf, export_txt
from test_reference_cases import inputs, run


class ViewStub:
    def __init__(self):self.options={};self.text='';self.rows={}
    def configure(self,**kw):self.options.update(kw)
    def delete(self,*items):
        self.text=''
        for item in items:self.rows.pop(item,None)
    def insert(self,*args,**kw):
        if 'values' in kw:self.rows[str(len(self.rows))]=kw['values']
        elif len(args)>1:self.text+=str(args[1])
    def get_children(self):return list(self.rows)


@pytest.fixture
def app(monkeypatch):
    interpreter=tk.Tcl()
    obj=gui.PuncoamentoApp.__new__(gui.PuncoamentoApp)
    obj.tk=interpreter.tk;obj._w='.';obj._loading=True;obj._redraw_job=None
    obj.after_idle=interpreter.after_idle;obj.after_cancel=interpreter.after_cancel
    obj._draw_scheme=lambda *a,**kw:None
    obj.vars={k:(tk.BooleanVar(master=interpreter,value=v) if isinstance(v,bool)
                 else tk.StringVar(master=interpreter,value=v)) for k,v in gui.DEFAULTS.items()}
    for name in ('var_resultado','var_status','var_rho','var_exemplo','var_steel'):
        setattr(obj,name,tk.StringVar(master=interpreter))
    for name in ('lbl_badge','txt_output','txt_diag','tree_summary','tree_steel'):setattr(obj,name,ViewStub())
    obj.last_verif=None;obj.last_report='';obj.dirty=True;obj.dialogs=[]
    for method in ('showerror','showinfo'):
        monkeypatch.setattr(gui.messagebox,method,lambda title,message,**kw:obj.dialogs.append((title,message)))
    for var in obj.vars.values():var.trace_add('write',obj._mark_dirty)
    obj._loading=False
    yield obj
    interpreter.eval('update idletasks')


@pytest.mark.parametrize('name,status',[
    ('Pilar interior - carga concêntrica, 300 kN','PASS_WITHOUT'),
    ('Pilar de bordo - excentricidade perpendicular interior','DETAIL_PENDING'),
    ('Pilar interior - armadura de punçoamento, 600 kN','DETAIL_PENDING'),
    ('Pilar interior - limite resistente com armadura','FAIL_MAX'),
    ('Pilar circular interior - esforço transverso, 450 kN','DETAIL_PENDING'),
    ('Pilar de canto - excentricidades interiores','DETAIL_PENDING'),
    ('Sapata circular - carga concêntrica e pressão uniforme','REQUIRES_REINFORCEMENT'),
    ('Pilar interior - setor ineficaz de 40 graus','BETA_PENDING'),
    ('Pilar interior - resistência máxima na face','FAIL_U0'),
    ('Pilar de bordo - g = 400 mm, VEd = 300 kN','BETA_PENDING'),
    ('Pilar de bordo - g = 400 mm, VEd = 400 kN','BETA_PENDING'),
    ('Pilar de bordo - g = 800 mm, VEd = 550 kN','BETA_PENDING'),
    ('Pilar de bordo - g = 400 mm, flexão biaxial','BETA_PENDING'),
    ('Pilar de bordo - g = 0 mm, VEd = 594 kN','DETAIL_PENDING'),
    ('Pilar de bordo - g = 343 mm, VEd = 594 kN','BETA_PENDING'),
    ('Pilar de bordo - g = 1 mm, VEd = 594 kN','BETA_PENDING'),
    ('Pilar de bordo - g = 0 mm, beta simplificado','DETAIL_PENDING'),
    ('Pilar de bordo - g = 343 mm, beta simplificado','DETAIL_PENDING'),
    ('Pilar interior - flexão biaxial, expressão (6.43)','PASS_WITHOUT'),
    ('Pilar interior - flexão uniaxial, expressão (6.39)','PASS_WITHOUT'),
    ('Pilar de bordo - excentricidade exterior uniaxial','DETAIL_PENDING'),
    ('Pilar de bordo - excentricidade exterior biaxial','DETAIL_PENDING'),
    ('Pilar de canto - excentricidade exterior biaxial','DETAIL_PENDING'),
    ('Pilar de canto - definição do modelo biaxial','BETA_PENDING'),
    ('Pilar interior - abertura quadrada a 600 mm','BETA_PENDING'),
    ('Pilar interior - abertura circular descentrada','BETA_PENDING'),
    ('Pilar interior - abertura retangular alongada','BETA_PENDING'),
    ('Pilar interior - abertura além de 6d','PASS_WITHOUT'),
    ('Pilar interior - duas aberturas ortogonais','BETA_PENDING'),])

def test_gui_examples_agree_with_engine(app,name,status):
    app.var_exemplo.set(name);app.carregar_exemplo();app.calcular()
    app.tk.eval('update idletasks')
    assert not app.dialogs
    assert app.last_verif.status==status
    assert not app.dirty
    assert app.var_resultado.get()==app.last_verif.snapshot()['conclusion']
    assert status in app.txt_output.text


def test_edit_invalidates_all_export_paths(app,tmp_path,monkeypatch):
    app.calcular();assert app._fresh_result() is not None
    old_hash=app.last_verif.snapshot()['input_hash']
    target=tmp_path/'stale.pdf'
    monkeypatch.setattr(gui.filedialog,'asksaveasfilename',lambda **kw:str(target))
    app.vars['V_Ed'].set('1000')
    assert app._fresh_result() is None
    app._export('pdf');assert not target.exists()
    with pytest.raises(ValueError):app._create_pdf(target)
    with pytest.raises(ValueError):app._build_professional_report_sections()
    app.calcular()
    assert app.last_verif.status=='FAIL_MAX'
    assert app.last_verif.snapshot()['input_hash']!=old_hash
    app._export('pdf');assert target.exists()


def test_invalid_recalculation_clears_old_pass(app):
    app.calcular();assert app.last_verif.status=='PASS_WITHOUT'
    app.vars['laje_d'].set('NaN');app.calcular()
    assert app.last_verif is None and app.dirty
    assert app.tree_summary.get_children()==[]
    assert 'VERIFICA' not in app.lbl_badge.options['text']
    assert app._fresh_result() is None


def test_depth_average_does_not_invalidate_fresh_calculation(app):
    app.vars['laje_dx'].set('0,19');app.vars['laje_dy'].set('0,23')
    app.calcular();app.tk.eval('update idletasks')
    assert not app.dialogs and not app.dirty
    assert app.last_verif.d==pytest.approx(.21)
    assert float(app.vars['laje_d'].get())==pytest.approx(.21)


def test_circular_json_save_load_round_trip(app,tmp_path,monkeypatch):
    app.vars['pilar_forma'].set('circular');app.vars['V_Ed'].set('321,5')
    app.vars['opening_sectors'].set('-20;20 | 100;125')
    target=tmp_path/'caso.json'
    monkeypatch.setattr(gui.filedialog,'asksaveasfilename',lambda **kw:str(target))
    monkeypatch.setattr(gui.filedialog,'askopenfilename',lambda **kw:str(target))
    app.guardar_caso();data=json.loads(target.read_text())
    assert data['V_Ed']==321500 and data['pilar_c2'] is None
    app.limpar();app.abrir_caso();app.calcular()
    assert not app.dialogs
    assert app.last_verif.snapshot()['inputs']['V_Ed']==321500
    assert app.last_verif.snapshot()['inputs']['opening_sectors']==[[-20,20],[100,125]]


@pytest.mark.parametrize('load,status',[(300000,'PASS_WITHOUT'),(600000,'DETAIL_PENDING'),(1000000,'FAIL_MAX'),(2000000,'FAIL_U0')])
def test_all_export_formats_retain_state_and_values(load,status,tmp_path):
    r=run(V_Ed=load,project='=2+2').snapshot()
    export_json(r,tmp_path/'result.json');export_txt(r,tmp_path/'result.txt')
    export_xlsx(r,tmp_path/'result.xlsx');export_pdf(r,tmp_path/'result.pdf')
    saved=json.loads((tmp_path/'result.json').read_text())
    assert saved['status']==status and saved['values']==r['values']
    assert status in (tmp_path/'result.txt').read_text()
    from openpyxl import load_workbook
    wb=load_workbook(tmp_path/'result.xlsx')
    assert wb['Resumo']['B2'].value==status
    entries=dict(wb['Entradas'].values)
    assert entries['VEd (N)']==load
    assert entries['Projeto']=="'=2+2"
    ws=wb['Verificacoes']
    for i,c in enumerate(r['checks'],2):
        assert ws.cell(i,2).value==pytest.approx(c['demand'])
        assert ws.cell(i,4).value==pytest.approx(c['utilization'])
        assert ws.cell(i,5).value==f'=IFERROR(B{i}/C{i},"")'
    # PDF parsing is an optional extra; the core test environment needs only
    # the published requirements. Visual QA is recorded separately.
    assert (tmp_path/'result.pdf').read_bytes().startswith(b'%PDF-')


@pytest.mark.parametrize('load,exit_code,status',[
    (300000,0,'PASS_WITHOUT'),(600000,1,'DETAIL_PENDING'),(1000000,1,'FAIL_MAX')])
def test_cli_exit_status_and_json_record(load,exit_code,status,tmp_path):
    source=tmp_path/'case.json';source.write_text(json.dumps(inputs(V_Ed=load)),encoding='utf-8')
    out=tmp_path/'out'
    command=[sys.executable,'-m','punching.cli',str(source),'--out',str(out)]
    completed=subprocess.run(command,capture_output=True,text=True,encoding='utf-8')
    assert completed.returncode==exit_code,completed.stderr
    assert json.loads((out/'resultado.json').read_text())['status']==status


def test_cli_rejects_invalid_input_without_success_record(tmp_path):
    source=tmp_path/'case.json';source.write_text(json.dumps(inputs(laje_d=0)),encoding='utf-8')
    out=tmp_path/'out'
    p=subprocess.run([sys.executable,'-m','punching.cli',str(source),'--out',str(out)],capture_output=True,text=True)
    assert p.returncode==2 and not out.exists()


def test_cli_records_pending_beta_with_nonzero_exit_code(tmp_path):
    source=Path('examples/15_bordo_343mm_594kN.json')
    out=tmp_path/'pending'
    p=subprocess.run([sys.executable,'-m','punching.cli',str(source),'--out',str(out)],capture_output=True,text=True)
    assert p.returncode==1
    r=json.loads((out/'resultado.json').read_text())
    assert r['status']=='BETA_PENDING' and r['values']['beta'] is None
    assert r['checks']==[] and r['reinforcement_rows']==[]


def test_gui_edge_distance_updates_geometry_inputs_and_invalidates(app,tmp_path,monkeypatch):
    app.var_exemplo.set('Pilar de bordo - g = 400 mm, VEd = 300 kN');app.carregar_exemplo();app.calcular()
    assert app.last_verif.u1==pytest.approx(3.706637061435917)
    app.vars['edge_distance_m'].set('0,80')
    assert app.dirty and app._fresh_result() is None
    app.calcular();assert app.last_verif.u1==pytest.approx(4.413274122871834)
    target=tmp_path/'edge_case.json'
    monkeypatch.setattr(gui.filedialog,'asksaveasfilename',lambda **kw:str(target))
    monkeypatch.setattr(gui.filedialog,'askopenfilename',lambda **kw:str(target))
    app.guardar_caso();app.limpar();app.abrir_caso();app.calcular()
    assert app.last_verif.edge_distance==pytest.approx(.8)
    assert app.last_verif.snapshot()['geometry']['column']['edge_distance_m']==pytest.approx(.8)


def test_gui_offset_rows_are_verified_before_anchorage_approval(app):
    app.var_exemplo.set('Pilar de bordo - g = 400 mm, VEd = 300 kN');app.carregar_exemplo()
    app.vars['V_Ed'].set('400');app.calcular()
    assert app.last_verif.status=='BETA_PENDING' and not app.tree_steel.rows
    app.vars['beta_mode'].set('simplificado');app.vars['simplified_applicable'].set(True);app.calcular()
    assert app.last_verif.status=='DETAIL_PENDING'
    assert len(app.tree_steel.rows)==len(app.last_verif.reinforcement_rows)>0
    app.vars['V_Ed'].set('400');app.vars['anchorage_confirmed'].set(True);app.calcular()
    assert app.last_verif.status=='PASS_WITH'
    assert app.lbl_badge.options['text']=='VERIFICA'
    assert all(c['passed'] for c in app.last_verif.detail_checks)


def test_gui_biaxial_edge_requires_beta_and_clears_invalid_result(app):
    app.var_exemplo.set('Pilar de bordo - g = 400 mm, flexão biaxial');app.carregar_exemplo();app.calcular()
    assert app.last_verif.status=='BETA_PENDING' and not app.tree_steel.rows
    app.vars['beta_mode'].set('manual');app.vars['beta_manual'].set('1,65')
    app.vars['beta_reference'].set('Ensaio da interface com fator prescrito; sem validação para projeto.');app.calcular()
    assert len(app.tree_steel.rows)==len(app.last_verif.reinforcement_rows)>0
    assert 'Ø10' in app.var_steel.get()
    assert app.tree_steel.rows['0'][2]==app.last_verif.reinforcement_rows[0]['n_legs']
    app.vars['laje_d'].set('0');app.calcular()
    assert not app.tree_steel.rows and app.last_verif is None


def test_gui_manual_beta_reference_roundtrip_and_pending_appearance(app,tmp_path,monkeypatch):
    app.var_exemplo.set('Pilar de bordo - g = 343 mm, VEd = 594 kN')
    app.carregar_exemplo();app.calcular()
    assert app.last_verif.status=='BETA_PENDING'
    assert app.lbl_badge.options['foreground']=='#956315'
    assert 'armadura não estão verificadas' in app.var_resultado.get()
    app.vars['beta_mode'].set('manual');app.vars['beta_manual'].set('1,40')
    app.vars['beta_reference'].set('Ensaio de persistência da fundamentação de beta.')
    app.calcular()
    assert app.last_verif.status=='DETAIL_PENDING'
    assert app.last_verif.beta==1.4
    target=tmp_path/'manual_case.json'
    monkeypatch.setattr(gui.filedialog,'asksaveasfilename',lambda **kw:str(target))
    monkeypatch.setattr(gui.filedialog,'askopenfilename',lambda **kw:str(target))
    app.guardar_caso();app.limpar();app.abrir_caso();app.calcular()
    assert app.last_verif.beta==1.4 and 'persistência' in app.last_verif.beta_terms['method']
    app.vars['beta_reference'].set('')
    assert app.dirty and app._fresh_result() is None
    app.calcular()
    assert app.last_verif is None and not app.tree_steel.rows
    assert app.dialogs and 'referência' in app.dialogs[-1][1]


def test_gui_general_biaxial_selection_invalidates_and_recalculates(app,tmp_path,monkeypatch):
    app.var_exemplo.set('Pilar de canto - definição do modelo biaxial')
    app.carregar_exemplo();app.calcular()
    assert app.last_verif.status=='BETA_PENDING'
    app.vars['allow_biaxial_envelope'].set(True)
    assert app.dirty and app._fresh_result() is None
    app.calcular()
    assert app.last_verif.status=='DETAIL_PENDING'
    assert 'hipótese' in app.txt_output.text or 'selecionada pelo projetista' in app.txt_output.text
    target=tmp_path/'biaxial.json'
    monkeypatch.setattr(gui.filedialog,'asksaveasfilename',lambda **kw:str(target))
    monkeypatch.setattr(gui.filedialog,'askopenfilename',lambda **kw:str(target))
    app.guardar_caso();app.limpar();app.abrir_caso();app.calcular()
    assert app.vars['allow_biaxial_envelope'].get() is True
    assert app.last_verif.beta_terms['basis']=='conservative_combination'


def test_gui_interior_method_is_preserved_in_saved_case(app,tmp_path,monkeypatch):
    app.var_exemplo.set('Pilar interior - flexão uniaxial, expressão (6.39)')
    app.carregar_exemplo();app.calcular()
    old=app.last_verif.beta
    app.vars['interior_beta_method'].set('ec2_643');app.calcular()
    assert app.last_verif.beta_terms['equation']=='6.43' and app.last_verif.beta!=old
    target=tmp_path/'interior.json'
    monkeypatch.setattr(gui.filedialog,'asksaveasfilename',lambda **kw:str(target))
    monkeypatch.setattr(gui.filedialog,'askopenfilename',lambda **kw:str(target))
    app.guardar_caso();app.limpar();app.abrir_caso();app.calcular()
    assert app.last_verif.interior_beta_method=='ec2_643'
