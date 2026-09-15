"""Independent arithmetic examples, mode isolation, orientation and persistence."""
from copy import deepcopy
import json
import math
import pytest
from punching.longitudinal import (derive,refresh_draft,draft_from_layout,layout_from_draft,
                                  DEFAULTS as LONG_DEFAULTS,DERIVED_KEYS,SPACINGS)
from punching.input_data import collect_inputs
from punching.core import PuncoamentoEC2
from punching.connections import ConnectionBook
from punching.collection_workflow import calculate_case,cached_calculation,rotate_joint
from punching.batch_workflow import bulk_edit,save_template,seed_case,assign_group
from punching.collection_reports import build_report,export_collection_pdf
from punching.reports import text_report,export_xlsx,export_pdf
from Punching_EC2_GUI import DEFAULTS
from test_batch_workflow import many,ready


def layout(base=10,extra=12,spacing=200):
    return dict(h_mm=250,cover_mm=30,outer_axis='x',layer_gap_mm=0,
                x=dict(base_mm=base,extra_mm=extra,spacing_mm=spacing),
                y=dict(base_mm=base,extra_mm=extra,spacing_mm=spacing))


def auto_draft(**kw):
    draft=deepcopy(DEFAULTS);draft.update(draft_from_layout(layout(**kw)))
    return draft


@pytest.mark.parametrize('extra,area,dx,dy,d',[
    (10,7.853981634,.215,.205,.210),(12,9.581857593,.214,.202,.208),
    (16,13.980087308,.212,.196,.204),(20,19.634954085,.210,.190,.200),
    (25,28.470683423,.2075,.1825,.195),(32,44.139376784,.204,.172,.188)])
def test_user_combinations_mm_to_cm2_per_m_and_layer_depths(extra,area,dx,dy,d):
    trace=derive(layout(extra=extra));v=trace['derived']
    assert v['laje_As_lx_cm2pm']==pytest.approx(area,abs=1e-8)
    assert v['laje_As_ly_cm2pm']==v['laje_As_lx_cm2pm']
    assert (v['laje_dx'],v['laje_dy'],v['laje_d'])==pytest.approx((dx,dy,d))
    assert trace['axes']['x']['pitch_mm']==100
    assert trace['axes']['x']['equivalent']==('Ø10 // 100' if extra==10 else '')


def test_equal_bars_alternated_equal_single_mesh_at_half_spacing():
    a=derive(layout(extra=10));b=derive(layout(extra=0,spacing=100))
    assert a['derived']==b['derived']
    assert a['axes']['x']['clear_mm']==b['axes']['x']['clear_mm']==90


@pytest.mark.parametrize('spacing',SPACINGS)
def test_all_practical_pitches_available_for_a_single_mesh(spacing):
    t=derive(layout(base=12,extra=0,spacing=spacing))
    assert t['axes']['x']['area_cm2pm']==pytest.approx(113.097335529*10/spacing)


def test_asymmetric_layers_and_nonzero_clear_gap():
    l=layout();l['y']=dict(base_mm=16,extra_mm=0,spacing_mm=200);l['layer_gap_mm']=5
    a=derive(l);l['outer_axis']='y';b=derive(l)
    assert [a['derived'][k] for k in ('laje_dx','laje_dy','laje_d')]==pytest.approx([.214,.195,.2045])
    assert [b['derived'][k] for k in ('laje_dx','laje_dy','laje_d')]==pytest.approx([.193,.212,.2025])
    assert a['derived']['laje_As_lx_cm2pm']==b['derived']['laje_As_lx_cm2pm']


@pytest.mark.parametrize('spacing',[100,125,150,175,190,225])
def test_nonstandard_intercalation_rejected_without_silent_rounding(spacing):
    with pytest.raises(ValueError,match='prátic'):derive(layout(spacing=spacing))


@pytest.mark.parametrize('name,value',[('h_mm',''),('cover_mm',''),('h_mm',float('nan')),('h_mm',True),('cover_mm',-1),('layer_gap_mm',-2),('outer_axis','z')])
def test_incomplete_or_invalid_physical_inputs_are_not_guessed(name,value):
    l=layout();l[name]=value
    with pytest.raises(ValueError):derive(l)


def test_bars_must_fit_and_clear_distance_is_checked():
    l=layout(extra=32);l['h_mm']=60
    with pytest.raises(ValueError,match='cabem'):derive(l)
    with pytest.raises(ValueError,match='distância livre'):derive(layout(base=40,extra=40),aggregate_mm=80)


def test_manual_mode_ignores_unfinished_hidden_layout_and_stays_numerically_identical():
    old=deepcopy(DEFAULTS);old.update(long_h_mm='?',long_x_base_mm='unfinished',long_mode='manual')
    inputs=collect_inputs(old,DEFAULTS)
    assert 'longitudinal_layout' not in inputs and inputs['laje_d']==.2
    v=PuncoamentoEC2(**inputs);v.verificar_puncoamento()
    assert 'longitudinal_trace' not in v.snapshot()


def test_automatic_mode_recomputes_stale_direct_areas_and_depths_in_gui_and_api():
    draft=auto_draft();draft.update({k:'stale' for k in DERIVED_KEYS})
    inputs=collect_inputs(draft,DEFAULTS);inputs.update(laje_d=9,laje_dx=9,laje_dy=9,laje_As_lx_cm2pm=999,laje_rho_l=.02)
    v=PuncoamentoEC2(**inputs);v.verificar_puncoamento();r=v.snapshot()
    assert r['values']['d']==pytest.approx(.208) and r['values']['dx']==pytest.approx(.214)
    assert r['inputs']['laje_As_lx_cm2pm']==pytest.approx(9.581857593)
    assert r['values']['rho_l']==pytest.approx(math.sqrt((9.581857593/2140)*(9.581857593/2020)))
    assert r['inputs']['laje_rho_l'] is None
    manual={k:v for k,v in r['inputs'].items() if k!='longitudinal_layout'}
    m=PuncoamentoEC2(**manual);m.verificar_puncoamento()
    assert m.snapshot()['values']==r['values'] and m.snapshot()['checks']==r['checks']
    assert 'Ø10 // 200 + Ø12 // 200' in text_report(r)


def test_bulk_auto_defaults_rotation_all_combos_reopen_and_only_calculated_report(tmp_path):
    b=many();ready(b);k=next(iter(b.cases));before=deepcopy(b.cases)
    l=layout();l['y']=dict(base_mm=16,extra_mm=0,spacing_mm=250)
    fields=draft_from_layout(l)
    assign_group(b,[k],'Teste automático')
    bulk_edit(b,[k],fields,DEFAULTS,only_empty=False,template=('group','Teste automático'))
    changed=[q for q,c in b.cases.items() if c['draft']['long_mode']=='automatic']
    assert len(changed)==3
    for q in set(b.cases)-set(changed):assert b.cases[q]==before[q]
    oldvals=refresh_draft(b.cases[k]['draft'])[0]
    rotate_joint(b,k,1,DEFAULTS)
    d=b.cases[k]['draft']
    assert d['long_outer_axis']=='y' and d['long_x_base_mm']=='16' and d['long_y_extra_mm']=='12'
    after,_=refresh_draft(d)
    assert after['laje_dx']==oldvals['laje_dy'] and after['laje_As_lx_cm2pm']==oldvals['laje_As_ly_cm2pm']
    result=calculate_case(b.cases[k],DEFAULTS);assert result['snapshot']
    report=build_report(b);assert len(report['entries'])==1
    reopened=ConnectionBook.from_payload(json.loads(json.dumps(b.payload())),DEFAULTS)
    assert cached_calculation(reopened.cases[k]) is not None
    assert sum(cached_calculation(c) is not None for c in reopened.cases.values())==1
    assert reopened.batch_settings==b.batch_settings
    export_collection_pdf(report,tmp_path/'collection.pdf')
    bulk_edit(b,[k],{'long_x_extra_mm':'20'},DEFAULTS,only_empty=False)
    assert not cached_calculation(b.cases[k])


def test_bulk_invalid_layout_is_atomic_and_cannot_edit_derived_values():
    b=many();ready(b);k=next(iter(b.cases))
    bulk_edit(b,[k],draft_from_layout(layout()),DEFAULTS,only_empty=False)
    before=deepcopy(b.payload())
    with pytest.raises(ValueError):bulk_edit(b,[k],{'long_h_mm':'55'},DEFAULTS,only_empty=False)
    assert b.payload()==before
    with pytest.raises(ValueError,match='modo automático'):bulk_edit(b,[k],{'laje_d':'.5'},DEFAULTS,only_empty=False)
    bulk_edit(b,[k],{'long_mode':'manual','laje_dx':'.23','laje_dy':'.24'},DEFAULTS,only_empty=False)
    assert b.cases[k]['draft']['laje_d']=='.235' or float(b.cases[k]['draft']['laje_d'])==.235


def test_legacy_drafts_preserved_as_manual_and_partial_new_layout_not_hidden():
    b=many();ready(b);payload=b.payload();original={}
    for k,c in payload['cases'].items():
        for field in LONG_DEFAULTS:c['draft'].pop(field)
        original[k]=deepcopy(c['draft'])
    restored=ConnectionBook.from_payload(payload,DEFAULTS)
    for k,c in restored.cases.items():
        assert c['draft']['long_mode']=='manual'
        assert {n:v for n,v in c['draft'].items() if n not in LONG_DEFAULTS}==original[k]
    damaged=b.payload();next(iter(damaged['cases'].values()))['draft'].pop('long_cover_mm')
    with pytest.raises(ValueError,match='incompatíveis'):ConnectionBook.from_payload(damaged,DEFAULTS)


def test_floor_template_recomputed_even_when_some_combinations_lack_forces():
    b=many();ready(b);k=next(iter(b.cases));floor=b.cases[k]['floor']
    save_template(b,'floor',floor,draft_from_layout(layout()))
    for c in b.cases.values():
        c['user_edited']=False;c['draft']=seed_case(b,c,DEFAULTS)
        calculate_case(c,DEFAULTS)
    ConnectionBook.from_payload(b.payload(),DEFAULTS)


def test_reports_preserve_layout_calculation_and_api_roundtrip(tmp_path):
    v=PuncoamentoEC2(**collect_inputs(auto_draft(),DEFAULTS));v.verificar_puncoamento();r=v.snapshot()
    v2=PuncoamentoEC2(**json.loads(json.dumps(r['inputs'])));v2.verificar_puncoamento()
    assert r['longitudinal_trace']==v2.snapshot()['longitudinal_trace']
    assert collect_inputs({**DEFAULTS,**draft_from_layout(r['inputs']['longitudinal_layout'])},DEFAULTS)['longitudinal_layout']==r['inputs']['longitudinal_layout']
    export_pdf(r,tmp_path/'automatic.pdf');export_xlsx(r,tmp_path/'automatic.xlsx')
    from openpyxl import load_workbook
    wb=load_workbook(tmp_path/'automatic.xlsx');assert 'ArmaduraLongitudinal' in wb.sheetnames
    assert any('dx = ' in str(row[0].value) for row in wb['ArmaduraLongitudinal'])
