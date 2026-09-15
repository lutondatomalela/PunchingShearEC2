"""Finite distance to one free edge: independent geometry and load checks."""
import json
import math
import pytest
from punching import geometry as geo
from punching.core import PuncoamentoEC2
from punching.reports import export_pdf,export_xlsx,text_report
from test_reference_cases import inputs,run


def p157(**changes):
    p=inputs(pilar_tipo='bordo',pilar_c1=.25,pilar_c2=.70,edge_distance_m=.4, beta_mode='manual',beta_manual=1.4,beta_reference='Ensaio numérico de distribuição com beta prescrito; não é validação do fator em projeto.')
    p.update(changes)
    v=PuncoamentoEC2(**p);v.verificar_puncoamento();return v


def test_edge_lengths_at_2d_and_conservative_face_bound():
    v=p157();r=v.snapshot()
    assert v.u0==pytest.approx(.85)
    assert v.u1==pytest.approx(.25+2*(.70+.40)+2*math.pi*.20)
    assert v.u1_eff==pytest.approx(3.706637061435917)
    comparison={c['kind']:c for c in r['contour_comparison']}
    assert comparison['fechado']['geometric_length']==pytest.approx(4.413274122871834)
    assert comparison['fechado']['admissible'] is False
    assert comparison['aberto_bordo']['selected'] is True
    assert v.u1_star is None and 'u1_star' not in r['geometry']
    for path in r['geometry']['u1']:
        assert min(y for x,y in path)>=-.75-1e-10
    points=r['geometry']['u1'][0]
    assert points[0][1]==pytest.approx(-.75)
    assert points[-1][1]==pytest.approx(-.75)


@pytest.mark.parametrize('gap,a',[(.4,.1),(.4,.2),(.4,.4),(.8,.4),(1.2,.5)])
def test_minimum_open_or_closed_contour_independent_formula(gap,a):
    c1=.25;c2=.7
    chosen,candidates=geo.select_edge_contour(c1,c2,a,gap)
    closed=2*(c1+c2)+2*math.pi*a
    opened=c1+2*(c2+gap)+math.pi*a
    expected=min(closed,opened) if gap>a+geo.EPS else opened
    assert chosen['effective_length']==pytest.approx(expected,abs=1e-12)


def test_setback_centroid_and_W_remain_geometric_quantities():
    cp=.25;cn=.70;gap=.40;a=.40;hx=cp/2;hy=cn/2
    u=cp+2*(cn+gap)+math.pi*a
    sy=-gap*(cn+gap)+cp*(hy+a)+math.pi*a*hy+2*a*a
    cy=sy/u
    wx=(cy+hy+gap)**2+(hy-cy)**2+cp*(hy+a-cy)+math.pi*a*(hy-cy)+2*a*a
    wy=2*(cn+gap)*(hx+a)+cp*cp/4+math.pi*a*hx+2*a*a
    v=p157(M_Edx=-25000,M_Edy=35000,edge_perp_interior=False)
    terms=v.beta_terms['candidates'][0]
    assert terms['cy']==pytest.approx(cy,abs=1e-12)
    assert terms['Wx']==pytest.approx(wx,abs=1e-12)
    assert terms['Wy']==pytest.approx(wy,abs=1e-12)
    assert v.beta==1.4
    assert 'ecx' not in terms and 'ecy' not in terms
    assert v.v_Ed_u1==pytest.approx(1.4*300000/u/.2/1e6)


def test_same_prescribed_beta_selects_shortest_admissible_contour():
    v=p157(edge_distance_m=.8)
    assert v.beta_terms['selected_contour']=='fechado'
    assert v.beta_terms['governing_contour']=='fechado'
    assert all(c['beta']==1.4 for c in v.beta_terms['candidates'])
    expected=1.4*v.V_Ed/v.u1_eff/v.d/1e6
    assert v.v_Ed_u1==pytest.approx(expected)


def test_zero_gap_retains_previous_flush_edge_method():
    v=run(pilar_tipo='bordo',M_Edx=30000,edge_distance_m=0)
    assert v.u0==pytest.approx(.4+2*min(.4,1.5*.2))
    assert v.u1==pytest.approx(2.456637061435917)
    assert v.beta==pytest.approx(1.1944922648241716)
    assert v.status=='DETAIL_PENDING'


def test_distant_edge_geometry_recovers_closed_contour_without_automatic_beta():
    v=p157(edge_distance_m=20,beta_mode='ec2',M_Edx=7000,M_Edy=23000)
    reference=run(pilar_c1=.25,pilar_c2=.7,M_Edx=7000,M_Edy=23000)
    assert v.u1==pytest.approx(reference.u1)
    assert v.beta is None and v.status=='BETA_PENDING'
    assert v.u0==pytest.approx(.85) # explicit edge capacity bound remains conservative


def test_distant_edge_closed_rings_still_verified_automatically():
    v=p157(edge_distance_m=20,V_Ed=450000,anchorage_confirmed=True)
    assert v.status=='PASS_WITH'
    assert v.geometry['u1_contour']==v.geometry['outer_contour']=='fechado'
    assert v.reinforcement_rows and all(r['passed'] for r in v.reinforcement_rows)


def test_closed_to_open_geometry_transition_before_outer_contour():
    switch=(2*.4-.25)/math.pi
    first,_=geo.select_edge_contour(.25,.7,switch-1e-6,.4)
    second,_=geo.select_edge_contour(.25,.7,switch+1e-6,.4)
    assert first['kind']=='fechado' and second['kind']=='aberto_bordo'
    assert abs(first['effective_length']-second['effective_length'])<2e-5
    v=p157(V_Ed=400000,anchorage_confirmed=True)
    assert v.status=='PASS_WITH'
    assert v.r_out_face>=2*v.d
    assert v.geometry['outer_utilization']<=1+1e-9
    assert v.geometry['outer_contour']=='aberto_bordo'
    assert v.reinforcement_rows and all(c['passed'] for c in v.detail_checks)
    assert v.Asw_por_perimetro>0
    assert v.Asw_sr_req>0
    assert all(v.Asw_sr_req>=r['Asw_sr_req_m2pm']-1e-12 for r in v.reinforcement_demands)


def test_moment_envelope_continuous_when_shortest_contour_switches():
    switch_gap=.25/2+math.pi*.4/2
    a=p157(edge_distance_m=switch_gap-1e-7,M_Edx=10000,M_Edy=8000)
    b=p157(edge_distance_m=switch_gap+1e-7,M_Edx=10000,M_Edy=8000)
    assert a.beta_terms['selected_contour']=='aberto_bordo'
    assert b.beta_terms['selected_contour']=='fechado'
    assert a.v_Ed_u1==pytest.approx(b.v_Ed_u1,abs=1e-6)


def test_sectors_also_trim_offset_contours_and_outer_boundary():
    v=p157(V_Ed=300000,opening_sectors=[(20,35)])
    assert v.u1_eff<v.u1
    assert v.contour_comparison[1]['effective_length']<v.contour_comparison[1]['geometric_length']
    assert v.status not in ('INVALID','ERROR')
    assert all(y>=-.75-1e-10 for path in v.geometry['u1'] for x,y in path)


def test_degenerate_offset_perimeter_is_invalid_and_reportable():
    v=p157(opening_sectors=[(1,359)])
    assert v.status=='INVALID'
    assert 'DADOS INVÁLIDOS' in text_report(v.snapshot())
    assert v.v_Ed_u1 is None


@pytest.mark.parametrize('changes',[
    {'edge_distance_m':-.01},{'edge_distance_m':float('nan')},
    {'edge_distance_m':float('inf')},{'edge_distance_m':True},
    {'edge_distance_m':101},{'edge_distance_m':.4,'pilar_tipo':'interior'},
    {'edge_distance_m':.4,'pilar_tipo':'canto'},
    {'edge_distance_m':.4,'pilar_forma':'circular'},
    {'edge_distance_m':.4,'is_sapata':True},
])
def test_reject_inconsistent_edge_inputs(changes):
    p=inputs(pilar_tipo='bordo');p.update(changes)
    with pytest.raises(ValueError):PuncoamentoEC2(**p)


def test_offset_detail_exports_actual_contributing_steel(tmp_path):
    v=p157(V_Ed=400000,anchorage_confirmed=True);r=v.snapshot()
    assert r['status']=='PASS_WITH'
    text=text_report(r)
    assert 'Conferência da distribuição proposta' in text
    assert 'Asw total' in text and 'ramos contribuem' in text
    export_pdf(r,tmp_path/'p157.pdf');export_xlsx(r,tmp_path/'p157.xlsx')
    from openpyxl import load_workbook
    wb=load_workbook(tmp_path/'p157.xlsx')
    assert wb['Resumo']['B2'].value=='PASS_WITH'
    assert wb['Contornos']['C3'].value==pytest.approx(3.706637061435917)
    assert wb['Fiadas'].max_row==len(r['reinforcement_rows'])+1
    assert wb['FiadasContornos'].max_row==len(r['reinforcement_rows'])+1
    assert wb['Ramos'].max_row==sum(row['n_legs'] for row in r['reinforcement_rows'])+1
    assert wb['Pormenorizacao'].max_row==len(r['detail_checks'])+1
    assert wb['ArmaduraNecessaria'].max_row>=2
