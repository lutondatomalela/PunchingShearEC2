"""Regression of the 594 kN connection and of explicit beta applicability.

Values with a prescribed beta validate arithmetic and geometry, not the
engineering justification of a beta input for a particular project.
"""
import json
import math
from pathlib import Path
import pytest
from punching.core import PuncoamentoEC2
from punching.reports import export_xlsx,export_json,export_pdf,text_report


def connection(gap=0.,**changes):
    p=dict(laje_d=.352,betão_fck=30,aço_fyk=500,aço_fywk=500,
           pilar_tipo='bordo',pilar_forma='retangular',pilar_c1=.7,pilar_c2=.25,
           edge_distance_m=gap,V_Ed=594000,M_Edx=-23000,M_Edy=-64000,
           laje_As_lx_cm2pm=13.98,laje_As_ly_cm2pm=13.98,
           edge_perp_interior=True,anchorage_confirmed=True)
    p.update(changes)
    v=PuncoamentoEC2(**p);v.verificar_puncoamento();return v


def test_flush_connection_against_644_and_actual_911_minimum():
    v=connection();d=.352;parallel=.7;normal=.25
    u=parallel+2*normal+2*math.pi*d
    reduced=parallel+2*min(normal/2,1.5*d)+2*math.pi*d
    W=parallel**2/4+normal*parallel+4*normal*d+8*d*d+math.pi*d*parallel
    beta=u/reduced+.45*u/W*(64000/594000)
    assert v.beta==pytest.approx(beta,rel=1e-12)
    assert v.status=='PASS_WITH' and v.u0==pytest.approx(1.2)
    assert v.v_Ed_u1==pytest.approx(.5676167056732232)
    assert len(v.reinforcement_rows)==2
    assert v.Asw_sr_calc*1e4==pytest.approx(13.936800412191)
    assert v.Asw_sr_min_u1_reference*1e4==pytest.approx(19.932317522095)
    for row in v.reinforcement_rows:
        # A row may be shorter than u1; local (9.11) is the relevant minimum.
        actual=math.pi*.01**2/4
        minimum=.08*math.sqrt(30)/500*row['sr']*row['st']/1.5
        assert row['Asw_m2']/row['sr']<v.Asw_sr_min_u1_reference
        assert actual>=minimum
        assert row['Asw_min_branch_m2']==pytest.approx(minimum)
        assert row['minimum_branch_utilization']==pytest.approx(minimum/actual)
        assert row['Asw_resistance_required_m2']==pytest.approx(v.Asw_sr_calc*row['sr'])
    report=text_report(v.snapshot())
    assert 'referência para o mínimo em u1' in report
    assert 'Mínimo local (9.11)' in report


@pytest.mark.parametrize('gap',[1e-12,1e-9,.001,.343,.352,.704,1.,20.])
def test_setback_ec2_never_manufactures_pass_or_failure(gap):
    v=connection(gap)
    assert v.status=='BETA_PENDING'
    assert v.beta is None and v.v_Ed_u1 is None and v.v_Ed_u0 is None
    assert v.Asw_sr_calc is None and v.armadura_necessaria is None
    assert not v.checks and not v.reinforcement_rows
    assert v.u1_eff>0 and v.v_Rd_c>0
    assert v.u0==pytest.approx(1.2)
    assert v.snapshot()['u0_terms']['conservative_edge_bound'] is True


@pytest.mark.parametrize('moment',[0.,1e-9,1.,23000.])
def test_zero_normal_moment_is_limit_of_inward_branch(moment):
    v=connection(M_Edx=moment)
    assert v.beta==pytest.approx(1.1475717093432567)
    assert v.status=='PASS_WITH'


@pytest.mark.parametrize('mode',['manual','simplificado'])
def test_fixed_model_is_continuous_at_zero_and_reduces_stress_with_gap(mode):
    options=dict(beta_mode=mode,simplified_applicable=True)
    if mode=='manual':options.update(beta_manual=1.4,beta_reference='Ensaio de continuidade com fator prescrito.')
    values=[connection(g,**options) for g in [0,1e-9,.001,.343,.704,1.5]]
    for v in values:
        closed=2*(.7+.25)+4*math.pi*.352
        opened=.7+2*(.25+v.edge_distance)+2*math.pi*.352
        expected=min(closed,opened) if v.edge_distance>2*.352 else opened
        assert v.u1_eff==pytest.approx(expected)
        assert v.beta==1.4 and v.u0==pytest.approx(1.2)
        assert v.v_Ed_u1==pytest.approx(1.4*594000/expected/.352/1e6)
        assert v.status in ('PASS_WITH','PASS_WITHOUT')
    assert values[0].v_Ed_u1==pytest.approx(values[1].v_Ed_u1,abs=1e-9)
    assert all(a.v_Ed_u1>=b.v_Ed_u1 for a,b in zip(values,values[1:]))


def test_simplified_conditions_cannot_be_inferred_or_bypassed_by_anchorage():
    with pytest.raises(ValueError,match='condições'):
        connection(.343,beta_mode='simplificado')
    for flag in (True,False):
        assert connection(.343,anchorage_confirmed=flag).status=='BETA_PENDING'
    v=connection(.343,beta_mode='simplificado',simplified_applicable=True)
    assert v.beta==1.4 and v.v_Ed_u1==pytest.approx(.5765455799205119)
    assert v.status=='PASS_WITH'


@pytest.mark.parametrize('changes',[
    {'beta_manual':None,'beta_reference':'Análise'},
    {'beta_manual':.99,'beta_reference':'Análise'},
    {'beta_manual':True,'beta_reference':'Análise'},
    {'beta_manual':float('nan'),'beta_reference':'Análise'},
    {'beta_manual':float('inf'),'beta_reference':'Análise'},
    {'beta_manual':1.4,'beta_reference':''},
    {'beta_manual':1.4,'beta_reference':'   '},
    {'beta_manual':1.4,'beta_reference':True},
])
def test_manual_beta_requires_finite_value_and_traceable_reference(changes):
    with pytest.raises(ValueError):connection(.343,beta_mode='manual',**changes)


def test_pending_export_is_not_a_calculation_failure_or_zero_steel(tmp_path):
    r=connection(.343).snapshot()
    export_json(r,tmp_path/'pending.json');export_xlsx(r,tmp_path/'pending.xlsx');export_pdf(r,tmp_path/'pending.pdf')
    saved=json.loads((tmp_path/'pending.json').read_text(encoding='utf-8'))
    assert saved['status']=='BETA_PENDING' and saved['values']['beta'] is None
    assert 'Nenhuma verificação concluída' in text_report(r)
    from openpyxl import load_workbook
    wb=load_workbook(tmp_path/'pending.xlsx')
    assert wb['Resumo']['B2'].value=='BETA_PENDING'
    assert wb['Verificacoes'].max_row==1 and wb['Fiadas'].max_row==1


def test_example_catalogue_uses_technical_descriptions_and_no_implicit_approvals():
    directory=Path(__file__).resolve().parents[1]/'examples'
    catalog=json.loads((directory/'catalog.json').read_text(encoding='utf-8'))
    assert len({x['title'] for x in catalog})==len(catalog)==30
    for entry in catalog:
        text=(entry['title']+' '+entry['description']).lower()
        assert 'auditoria' not in text and 'imagem' not in text
        p=json.loads((directory/entry['file']).read_text(encoding='utf-8'))
        assert p['anchorage_confirmed'] is False
        assert p['support']==entry['title']
        if p['beta_mode']=='simplificado':
            assert 'pressupõe' in entry['description'] and '25%' in entry['description']
